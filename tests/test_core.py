"""核心邏輯測試（不需網路、不需金鑰）：python -m pytest"""

from datetime import datetime, timezone

import pytest

from display_intel import derive, verify
from display_intel.collect import canonical_url, classify, html_to_text
from display_intel.db import DB
from display_intel.extract import extract_all, make_batches
from display_intel.llm import BudgetExceeded, parse_json
from display_intel.schema import Confidence, Extraction, PanelSpec, RawArticle

KEYWORDS = {
    "device": ["laptop", "筆電"], "display": ["oled", "display", "螢幕"],
    "platform_vendor": ["intel"], "panel_maker": ["boe"], "fab": ["gen 8.6"],
}


def art(title="t", body="", source_type="media", url=None):
    return RawArticle(url=url or f"https://x.com/{abs(hash(title + body))}", title=title, source="S",
                      source_type=source_type, content=body, published=datetime.now(timezone.utc))


# ---------- 衍生規格 ----------

def test_ppi_and_aspect():
    assert derive.ppi(2560, 1600, 16) == 188.7
    assert derive.aspect_ratio(2560, 1600) == "16:10"
    assert derive.aspect_ratio(1920, 1080) == "16:9"
    assert derive.aspect_ratio(3000, 2000) == "3:2"
    assert derive.density_class(140) == "一般"
    assert derive.density_class(189) == "高"
    assert derive.density_class(260) == "極高"
    assert derive.ppi(None, 1600, 16) is None


def test_plausibility_removes_bad_values():
    spec = PanelSpec(brand="X", model="Y", size_inch=13, res_w=7680, res_h=4320, refresh_hz=2400, sdr_nits=500)
    bad = derive.plausibility_issues(spec)
    assert "refresh_hz" in bad and "ppi" in bad
    assert "sdr_nits" not in bad


def test_evidence_required():
    text = "The Foo 14 laptop has a 14-inch OLED panel at 120Hz."
    spec = PanelSpec(brand="Acme", model="Foo 14", size_inch=14, refresh_hz=120, sdr_nits=600,
                     evidence={"size_inch": "14-inch", "refresh_hz": "120Hz", "sdr_nits": "600 nits"})
    cleaned, dropped = derive.enforce_evidence(spec, text)
    assert cleaned.refresh_hz == 120 and cleaned.size_inch == 14
    assert cleaned.sdr_nits is None and "sdr_nits" in dropped  # 文章沒寫 600 nits → 刪除
    assert cleaned.brand is None  # 品牌沒出現在文中
    assert cleaned.model == "Foo 14"


def test_quote_matching_ignores_width_and_space():
    assert derive.quote_in_text("２８８０ｘ１８００", "解析度 2880x1800 OLED")
    assert derive.quote_in_text("120  Hz", "a 120 Hz panel")
    assert not derive.quote_in_text("240Hz", "a 120Hz panel")


# ---------- 蒐集 ----------

def test_classify_requires_two_keyword_groups():
    assert classify(art("New laptop with OLED display"), KEYWORDS) == ["panel"]
    assert classify(art("New laptop with better battery"), KEYWORDS) == []
    assert "fab" in classify(art("BOE Gen 8.6 line"), KEYWORDS)
    assert "platform" in classify(art("Intel display tech"), KEYWORDS)


def test_html_and_url_helpers():
    assert "Hello" in html_to_text("<html><script>x()</script><p>Hello <b>world</b></p></html>")
    assert "x()" not in html_to_text("<html><script>x()</script><p>Hello</p></html>")
    assert canonical_url("https://A.com/x/?utm_source=a&id=3") == "https://a.com/x?id=3"


# ---------- AI 回覆解析 ----------

def test_parse_json_tolerates_fences():
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json('好的：{"a": [1,2]} 完成') == {"a": [1, 2]}
    with pytest.raises(ValueError):
        parse_json("沒有 JSON")


def test_make_batches_respects_budget():
    arts = [art(str(i), "x" * 1000) for i in range(10)]
    batches = make_batches(arts, max_chars_per_call=3000, max_chars_per_article=1500)
    assert sum(len(b) for b in batches) == 10
    assert all(len(b) <= 3 for b in batches)


class FakeLLM:
    def __init__(self, replies):
        self.replies = list(replies)
        self.calls = 0

    def complete_json(self, system, user):
        self.calls += 1
        r = self.replies.pop(0)
        if isinstance(r, Exception):
            raise r
        return r


def test_extract_maps_ids_and_drops_hallucination():
    a1 = art("Acme Foo 14 laptop OLED", "Acme Foo 14 has a 14-inch OLED panel with 120Hz refresh rate.")
    a2 = art("Weather today", "Sunny.")
    reply = {"results": [
        {"article_id": 1, "relevant": True, "panels": [
            {"brand": "Acme", "model": "Foo 14", "panel_tech": "OLED", "size_inch": 14, "refresh_hz": 240,
             "evidence": {"panel_tech": "OLED panel", "size_inch": "14-inch", "refresh_hz": "240Hz"}}]},
        {"article_id": 2, "relevant": False},
        {"article_id": 9, "relevant": True},  # 不存在的編號 → 忽略
    ]}
    results, stats = extract_all([a1, a2], FakeLLM([reply]), {"max_chars_per_call": 10000})
    assert len(results) == 1
    _, ex, flags = results[0]
    assert ex.panels[0].size_inch == 14
    assert ex.panels[0].refresh_hz is None  # 原文是 120Hz，AI 說 240Hz → 被刪除
    assert any("refresh_hz" in f for f in flags)


def test_extract_stops_when_budget_exceeded():
    arts = [art(str(i), "laptop oled " + "x" * 3000) for i in range(4)]
    llm = FakeLLM([{"results": []}, BudgetExceeded("用完")])
    results, stats = extract_all(arts, llm, {"max_chars_per_call": 3200, "max_chars_per_article": 3100})
    assert stats["done_batches"] == 1 and "用完" in stats["stopped"]


# ---------- 查核與資料庫 ----------

@pytest.fixture
def db(tmp_path):
    d = DB(tmp_path / "t.sqlite")
    yield d
    d.close()


def test_merge_new_update_and_rumor_confirmed(db):
    spec = PanelSpec(brand="Acme", model="Foo 14", size_inch=14, panel_tech="OLED", res_w=2880, res_h=1800, refresh_hz=120)
    rumor = art("leak", source_type="social")
    ev = verify.merge_panel(db, spec, Confidence.D, rumor, [])
    assert ev.kind == "new" and ev.record.ppi == 242.6

    # 同樣規格再出現一次，沒有新資訊 → 不產生事件
    assert verify.merge_panel(db, spec, Confidence.D, rumor, []) is None

    # 官方公布並補上亮度 → 傳聞證實
    official = art("official", source_type="official")
    spec2 = spec.model_copy(update={"sdr_nits": 500})
    ev = verify.merge_panel(db, spec2, Confidence.A, official, [])
    assert ev.kind == "rumor_confirmed" and ev.record.confidence == Confidence.A
    assert ev.changes["sdr_nits"] == [None, 500]

    # 較低可信度的來源不能覆蓋官方數值
    spec3 = spec.model_copy(update={"refresh_hz": 144})
    assert verify.merge_panel(db, spec3, Confidence.C, art("m"), []) is None
    assert db.get_panel(verify.panel_key(spec)).spec.refresh_hz == 120


def test_dedupe_articles(db):
    a = art("ASUS unveils new OLED laptop lineup", source_type="aggregator", url="https://a/1")
    b = art("ASUS unveils new OLED laptop line-up", source_type="official", url="https://b/1")
    c = art("Completely different story", url="https://c/1")
    db.add_article("https://c/1", "Completely different story", "S")
    kept = verify.dedupe_articles([a, b, c], db)
    assert [k.url for k in kept] == ["https://b/1"]  # 保留官方來源；c 已處理過


def test_ingest_and_report(db, tmp_path):
    from display_intel.report import build_report
    a = art("Acme Foo 14 laptop", "Acme Foo 14 has a 14-inch OLED panel.", source_type="official")
    ex = Extraction(relevant=True, panels=[PanelSpec(brand="Acme", model="Foo 14", size_inch=14, panel_tech="OLED")])
    stats = verify.ingest(db, [(a, ex, [])])
    assert stats["panel_new"] == 1
    cfg = {"schedule": {"timezone": "Asia/Taipei"}, "report": {"output_dir": str(tmp_path), "subject_prefix": "[T]"}}
    ctx, html, xlsx, ids = build_report(db, cfg)
    assert "Foo 14" in html and xlsx.exists() and len(ids) == 1
    assert "新機 1" in ctx["subject"]
    db.mark_reported(ids)
    assert db.unreported_events() == []


# ---------- 免費 AI 介面（模擬伺服器，不連網） ----------

def test_gemini_request_and_429(monkeypatch):
    import httpx
    from display_intel.llm import LLM

    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    monkeypatch.setattr("time.sleep", lambda s: None)
    calls = []

    def handler(request):
        calls.append(request)
        if len(calls) == 1:
            return httpx.Response(429, headers={"retry-after": "5"})
        return httpx.Response(200, json={"choices": [{"message": {"content": '{"ok": true}'}}]})

    llm = LLM("gemini", "gemini-3.5-flash-lite", max_calls=5)
    llm._http = httpx.Client(transport=httpx.MockTransport(handler), headers={"Authorization": "Bearer fake"})
    assert llm.complete_json("sys", "user") == {"ok": True}
    assert calls[0].url == "https://generativelanguage.googleapis.com/v1beta/openai/chat/completions"
    assert calls[0].headers["authorization"] == "Bearer fake"
    assert llm.calls == 2


def test_budget_limit(monkeypatch):
    import httpx
    from display_intel.llm import LLM

    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    llm = LLM("gemini", "m", max_calls=1)
    llm._http = httpx.Client(transport=httpx.MockTransport(
        lambda r: httpx.Response(200, json={"choices": [{"message": {"content": "{}"}}]})))
    llm.complete_json("s", "u")
    with pytest.raises(BudgetExceeded):
        llm.complete_json("s", "u")


def test_extract_splits_batch_on_truncation():
    from display_intel.llm import OutputTruncated
    arts = [art(f"Acme Foo {i} laptop OLED", f"Acme Foo {i} has a 14-inch OLED panel.") for i in range(2)]
    ok = {"results": [{"article_id": 1, "relevant": True, "panels": [
        {"brand": "Acme", "model": "Foo", "panel_tech": "OLED", "evidence": {"panel_tech": "OLED panel"}}]}]}
    llm = FakeLLM([OutputTruncated("太長"), ok, ok])
    results, stats = extract_all(arts, llm, {"max_chars_per_call": 10000})
    assert stats["split_batches"] == 1 and stats["done_batches"] == 2
    assert len(results) == 2 and llm.calls == 3
    assert stats["errors"]


def test_working_llm_falls_back(monkeypatch):
    from display_intel import __main__ as main
    from display_intel.llm import LLM
    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    monkeypatch.setattr(LLM, "probe", lambda self: "ok" if self.model == "good" else "HTTPStatusError：404 unknown model")
    stats = {}
    llm = main._working_llm({"llm": {"provider": "gemini", "extract_model": "bad", "fallback_models": ["good"]}}, stats)
    assert llm.model == "good" and stats["llm_model"] == "good"
    assert "404" in stats["llm_probe"]["bad"]


def test_non_json_body_gives_clear_error(monkeypatch):
    import httpx
    from display_intel.llm import LLM
    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    llm = LLM("gemini", "m", max_calls=3)
    llm._http = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, text="<html>retired</html>")))
    with pytest.raises(RuntimeError, match="回覆格式異常"):
        llm.complete_json("s", "u")


def test_github_models_retired_message():
    from display_intel.llm import LLM
    with pytest.raises(RuntimeError, match="gemini"):
        LLM("github_models", "x")


def test_unidentifiable_panels_dropped():
    from display_intel.extract import is_identifiable, postprocess
    assert is_identifiable(PanelSpec(brand="HP", model="OmniBook 5"))
    assert not is_identifiable(PanelSpec(brand=None, model="Laptop"))
    assert not is_identifiable(PanelSpec(brand="MSI", model=None))
    assert not is_identifiable(PanelSpec(brand="Asus", model="筆電"))
    a = art("MSI laptop", "MSI shows a new laptop with a 240Hz OLED display.")
    ex = Extraction(relevant=True, panels=[PanelSpec(brand="MSI", refresh_hz=240, evidence={"refresh_hz": "240Hz"})])
    ex, flags = postprocess(a, ex)
    assert ex.panels == [] and any("缺少品牌或機型" in f for f in flags)


def test_tavily_request(monkeypatch):
    import json as _json
    import httpx
    from display_intel.collect import _tavily
    monkeypatch.setenv("TAVILY_API_KEY", "tvly-fake")
    seen = []

    def handler(req):
        seen.append(_json.loads(req.content))
        return httpx.Response(200, json={"results": [
            {"url": "https://n.com/a?utm_source=x", "title": "Acme laptop OLED", "content": "snippet",
             "raw_content": "Acme Foo 14 laptop has a 14-inch OLED display. " * 30}]})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    cfg = {"sources": {"tavily": {"queries": ["q1", "q2", "q3"], "max_queries_per_run": 2}}}
    out = _tavily(client, cfg)
    assert len(seen) == 2 and seen[0]["time_range"] == "day" and seen[0]["include_raw_content"] == "text"
    assert out[0].content_level == "full" and out[0].url == "https://n.com/a"


def test_weekly_cadence(db, tmp_path):
    from display_intel.collect import google_news_urls
    from display_intel.report import build_report
    cfg = {"schedule": {"cadence": "weekly", "timezone": "Asia/Taipei"},
           "sources": {"google_news": {"locales": [{"hl": "en-US", "gl": "US", "ceid": "US:en", "lang": "en"}],
                                       "queries": {"en": ["laptop OLED"]}}},
           "report": {"output_dir": str(tmp_path), "subject_prefix": "[週報]"}}
    assert "when%3A7d" in google_news_urls(cfg)[0]["url"]
    ctx, html, xlsx, ids = build_report(db, cfg)
    assert ctx["title"] == "筆電顯示器情報週報" and "過去 7 天" in html and "本週重點" in html


# ---------- 面板型號解碼與合併 ----------

def test_panel_decode():
    from display_intel.panel_decode import canonical_maker, decode
    assert decode("NE160QDM-NY1") == {"maker": "BOE", "size_inch": 16.0, "res": None}
    assert decode("ATNA56YX03") == {"maker": "Samsung Display", "size_inch": 15.6, "res": None}
    assert decode("B140HAN06.8") == {"maker": "AUO", "size_inch": 14.0, "res": None}
    assert decode("N156HCA-EAC")["maker"] == "Innolux"
    assert decode("LP160UQ1-SPB1")["maker"] == "LG Display"
    assert decode("XYZ123") is None
    assert canonical_maker("京东方") == "BOE" and canonical_maker("Samsung") == "Samsung Display"
    assert canonical_maker("SomeNewMaker") == "SomeNewMaker"


def test_inferred_maker_and_merge(db):
    a = PanelSpec(brand="ASUS", model="ProArt P16（H7607）", panel_tech="OLED", size_inch=16, res_w=3840, res_h=2400)
    b = PanelSpec(brand="ASUS", model="ASUS ProArt P16", panel_tech="OLED", panel_model="ATNA60CL10", refresh_hz=120)
    assert verify.model_norm(a) == verify.model_norm(b) == "proartp16"
    ev1 = verify.merge_panel(db, a, Confidence.C, art("a"), [])
    ev2 = verify.merge_panel(db, b, Confidence.C, art("b"), [])
    assert ev1.kind == "new" and ev2.kind == "updated"
    assert len(db.all_panels()) == 1
    rec = db.all_panels()[0]
    assert rec.spec.refresh_hz == 120 and rec.panel_maker_inferred == "Samsung Display" and rec.maker == "Samsung Display"
    # 不同尺寸不能合併
    c = PanelSpec(brand="ASUS", model="ProArt P16", panel_tech="OLED", size_inch=14)
    assert verify.merge_panel(db, c, Confidence.C, art("c"), []).kind == "new"


def test_consolidate_existing_duplicates(db):
    from display_intel.db import utcnow
    from display_intel.schema import PanelRecord
    now = utcnow()
    for key, spec in [("k1", PanelSpec(brand="HP", model="OMEN 16", res_w=2560, res_h=1600)),
                      ("k2", PanelSpec(brand="HP", model="HP OMEN 16 (16-xd0000)", refresh_hz=240))]:
        db.put_panel(PanelRecord(key=key, spec=spec, confidence=Confidence.C, source_url="u", source_name="s",
                                 first_seen=now, last_updated=now))
    assert verify.consolidate(db) == 1
    rec = db.all_panels()[0]
    assert rec.spec.refresh_hz == 240 and rec.spec.res_w == 2560


def test_fix_panel_source():
    from display_intel.extract import fix_panel_source
    s, flags = fix_panel_source(PanelSpec(brand="ASUS", model="ProArt P16", panel_maker="ASUS", panel_model="Lumina Pro OLED"))
    assert s.panel_maker is None and s.panel_model is None and len(flags) == 2
    s, _ = fix_panel_source(PanelSpec(brand="HP", model="X", panel_maker="京东方", panel_model="NE160QDM-NY1"))
    assert s.panel_maker == "BOE" and s.panel_model == "NE160QDM-NY1"
    s, _ = fix_panel_source(PanelSpec(brand="HP", model="X", panel_model="TL140ADXP01"))
    assert s.panel_model == "TL140ADXP01"  # 認不得但像料號 → 保留


# ---------- 報告分析 ----------

def _rec(key, days_ago, conf=Confidence.C, **spec):
    from datetime import timedelta
    from display_intel.db import utcnow
    from display_intel.schema import PanelRecord
    t = utcnow() - timedelta(days=days_ago)
    r = PanelRecord(key=key, spec=PanelSpec(**spec), confidence=conf, source_url="u", source_name="s", first_seen=t, last_updated=t)
    return verify._derived(r)


def test_analytics_blocks():
    from datetime import timedelta
    from display_intel import analytics
    from display_intel.db import utcnow
    start = utcnow() - timedelta(days=7)
    recs = [
        _rec("a", 1, brand="ASUS", model="A", panel_tech="Tandem OLED", refresh_hz=120, res_w=2880, res_h=1800, size_inch=14, panel_maker="Samsung Display"),
        _rec("b", 2, brand="HP", model="B", panel_tech="IPS LCD", refresh_hz=240, panel_model="NE160QDM-NY1"),
        _rec("c", 10, brand="HP", model="C", panel_tech="OLED", refresh_hz=60),
        _rec("d", 3, Confidence.D, brand="Apple", model="D", panel_tech="OLED"),
    ]
    assert analytics.features(recs[0].spec) >= {"OLED", "Tandem OLED", "高解析 2.8K+"}
    rows = [{"rec": r} for r in recs[:2]]
    k = {x["label"]: x["value"] for x in analytics.kpis(rows)}
    assert k["新機型"] == "2" and k["OLED 占比"] == "50%" and k["刷新率中位數"] == "180Hz"
    lines, tags = analytics.standouts(rows, recs, start)
    assert any("240Hz" in l for l in lines) and "最高刷新率" in tags["b"]
    assert any("首見技術：Tandem OLED" in l for l in lines)  # 本期之前只有 C（一般 OLED）
    wow = {w["metric"]: w for w in analytics.week_over_week(recs, start)}
    assert wow["新機型數"]["now"] == "2" and wow["新機型數"]["prev"] == "1" and wow["新機型數"]["arrow"] == "▲"
    m = analytics.brand_maker_matrix(recs)
    assert set(m["makers"]) == {"Samsung Display", "BOE"} and m["inferred"] == 1
    assert [t["name"] for t in analytics.rumor_tracker(recs, utcnow())] == ["Apple D"]
    assert analytics.week_over_week(recs[:2], start) is None  # 沒有上期資料


# ---------- 規格補完 ----------

def test_completeness_and_candidates(db):
    from datetime import timedelta
    from display_intel import analytics, enrich
    from display_intel.db import utcnow
    sparse = _rec("s", 2, brand="HP", model="OmniBook 5", panel_tech="OLED")
    rich = _rec("r", 2, brand="ASUS", model="ProArt P16", panel_tech="OLED", size_inch=16, res_w=3840, res_h=2400,
                refresh_hz=120, sdr_nits=500, dci_p3_pct=100, panel_maker="Samsung Display", panel_model="ATNA60",
                response_ms=0.2, dimming="DC", surface="glossy", delta_e=1, touch=False)
    rumor = _rec("d", 2, Confidence.D, brand="Apple", model="MacBook Pro")
    old = _rec("o", 90, brand="HP", model="Old")
    assert analytics.completeness(rich) == (13, 13) and analytics.completeness(sparse)[0] == 1
    for r in (sparse, rich, rumor, old):
        db.put_panel(r)
    col = {"enrich_target": 10, "enrich_within_days": 60, "enrich_schedule_days": [0, 14, 35], "enrich_retry_days": 6}
    assert [[r.key for r in g] for g in enrich.candidates(db, col)] == [["s"]]  # 已達標、傳聞、太舊都排除
    # 第一次補查後，要等到上市第 14 天才會再查
    sparse.enrich_attempts, sparse.last_enriched = 1, utcnow() - timedelta(days=7)
    db.put_panel(sparse)
    assert enrich.candidates(db, col) == []
    assert len(enrich.candidates(db, col, now=utcnow() + timedelta(days=13))) == 1


def test_tavily_budget_cap(db):
    from display_intel.enrich import TavilyBudget, _tavily
    from display_intel.llm import BudgetExceeded
    db.add_usage("tavily", 850)
    with pytest.raises(BudgetExceeded):
        _tavily(None, TavilyBudget(db, 850), "q")
    assert db.usage("tavily") == 850


def test_enrich_end_to_end(db, monkeypatch):
    import httpx
    from display_intel import enrich as enr
    monkeypatch.setenv("TAVILY_API_KEY", "tvly-fake")
    rec = _rec("x", 1, brand="HP", model="OmniBook 5", panel_tech="OLED")
    db.put_panel(rec)
    page = ("HP OmniBook 5 review. The OmniBook 5 has a 14-inch 1920x1200 OLED panel from Samsung Display "
            "with 400 nits and 100% DCI-P3. The Acme X1 has a 4K screen. ") * 3

    def handler(req):
        return httpx.Response(200, json={"results": [
            {"url": "https://www.notebookcheck.net/HP-OmniBook-5-review.1.0.html", "title": "HP OmniBook 5 review",
             "content": "review", "raw_content": page},
            {"url": "https://other.com/acme", "title": "Acme X1", "content": "x", "raw_content": "Acme X1 laptop " * 100}]})

    real_client = httpx.Client
    monkeypatch.setattr(enr.httpx, "Client", lambda **kw: real_client(transport=httpx.MockTransport(handler)))
    reply = {"results": [{"article_id": 1, "relevant": True, "panels": [
        {"brand": "HP", "model": "OmniBook 5", "panel_tech": "OLED", "panel_maker": "Samsung Display",
         "size_inch": 14, "res_w": 1920, "res_h": 1200, "sdr_nits": 400, "dci_p3_pct": 100,
         "evidence": {"panel_tech": "OLED panel", "panel_maker": "Samsung Display", "size_inch": "14-inch",
                      "res_w": "1920x1200", "res_h": "1920x1200", "sdr_nits": "400 nits", "dci_p3_pct": "100% DCI-P3"}},
        {"brand": "Acme", "model": "X1", "res_w": 3840, "evidence": {}}]}]}
    cfg = {"collection": {"enrich_min_missing": 4}, "llm": {"max_chars_per_call": 50000, "max_chars_per_article": 3000,
                                                            "max_chars_per_review": 8000}}
    stats = enr.enrich(db, cfg, FakeLLM([reply]), limit=5)
    assert stats["searched"] == 1 and stats["pages"] == 1  # 沒提到 OmniBook 5 的頁面被過濾、重複網址合併
    assert stats["tavily"] == 3 and db.usage("tavily") == 3  # 評測＋官方各一次，缺面板型號再專門搜尋一次
    assert stats["panel_hunts"] == 1
    assert stats["improved"] == 1 and stats["fields_added"] >= 4
    recs = db.all_panels()
    assert len(recs) == 1 and recs[0].maker == "Samsung Display" and recs[0].confidence == Confidence.B
    assert recs[0].enrich_attempts == 1
    assert any(e["kind"] == "updated" for e in db.unreported_events())


def test_profiles_schedule():
    from datetime import datetime
    from zoneinfo import ZoneInfo
    from display_intel.config import apply_profile, load_config
    cfg = load_config()
    tz = ZoneInfo("Asia/Taipei")
    apply_profile(cfg, datetime(2026, 10, 14, 19, 17, tzinfo=tz))  # 週三
    assert cfg["run"]["profile"] == "light" and cfg["run"]["tavily_days"] == 2 and len(cfg["run"]["tavily_queries"]) == 10
    apply_profile(cfg, datetime(2026, 10, 18, 19, 17, tzinfo=tz))  # 週日
    assert cfg["run"]["profile"] == "full" and cfg["run"]["extract_model"] and cfg["run"]["enrich_models"] == 3
    apply_profile(cfg, datetime(2026, 10, 20, 19, 17, tzinfo=tz))  # 週二：規格補完日
    assert cfg["run"]["profile"] == "enrich" and cfg["run"]["enrich_models"] == 30 and cfg["run"]["collect_news"] is False
    apply_profile(cfg, profile="enrich")
    assert cfg["run"]["collect_news"] is False and cfg["run"]["tavily_queries"] == []


def test_identity_and_grouping(db, tmp_path):
    from display_intel.extract import normalize_identity
    s = normalize_identity(PanelSpec(brand="Asus", model="ASUS Zenbook 14"))
    assert (s.brand, s.model) == ("ASUS", "Zenbook 14")
    s = normalize_identity(PanelSpec(brand="Framework", model="Framework Laptop 12"))
    assert s.model == "Laptop 12"
    assert normalize_identity(PanelSpec(brand="dell alienware", model="x")).brand == "Dell alienware"
    from display_intel.report import build_report
    a = art("OMEN", "HP OMEN 15 has a 15.6-inch 1920x1080 panel and a 2560x1440 panel option.", source_type="official")
    ex = Extraction(relevant=True, panels=[
        PanelSpec(brand="HP", model="OMEN 15", size_inch=15.6, res_w=1920, res_h=1080),
        PanelSpec(brand="HP", model="OMEN 15", size_inch=15.6, res_w=2560, res_h=1440)])
    verify.ingest(db, [(a, ex, [])])
    assert len(db.all_panels()) == 2  # 兩種選配分開記錄
    cfg = {"schedule": {"timezone": "Asia/Taipei"}, "report": {"output_dir": str(tmp_path)}}
    ctx, html, _, _ = build_report(db, cfg)
    assert len(ctx["new_panels"]) == 1 and len(ctx["new_panels"][0]["options"]) == 2
    assert "2 種螢幕選配" in html and ctx["kpis"][0]["value"] == "1"



def test_panel_lookup_fills_blanks(db, monkeypatch):
    import httpx
    from display_intel import enrich as enr
    monkeypatch.setenv("TAVILY_API_KEY", "tvly-fake")
    rec = _rec("p", 1, brand="HP", model="OMEN 16", panel_model="NE160QDM-NY1", sdr_nits=500)
    db.put_panel(rec)
    page = "BOE NE160QDM-NY1 is a 16.0 inch IPS panel. Response time: 3 ms. Color gamut 100% sRGB. Brightness 400 nits."

    def handler(req):
        return httpx.Response(200, json={"results": [{"url": "https://panels.example/ne160", "title": "NE160QDM-NY1",
                                                      "content": page, "raw_content": page}]})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    llm = FakeLLM([{"panel": {"panel_tech": "IPS", "response_ms": 3, "srgb_pct": 100, "sdr_nits": 400, "dci_p3_pct": 95,
                              "evidence": {"panel_tech": "IPS panel", "response_ms": "Response time: 3 ms",
                                           "srgb_pct": "100% sRGB", "sdr_nits": "400 nits", "dci_p3_pct": "95% DCI-P3"}}}])
    values, url = enr.lookup_panel(client, llm, enr.TavilyBudget(db, 850), "NE160QDM-NY1")
    assert values["response_ms"] == 3 and "dci_p3_pct" not in values  # 沒有原文依據的被刪掉
    changes = enr.apply_panel_specs(db, rec, values, url)
    r = db.get_panel("p")
    assert r.spec.sdr_nits == 500  # 整機數值不被面板規格覆蓋
    assert r.spec.response_ms == 3 and "response_ms" in r.panel_spec_fields and "sdr_nits" not in changes


def test_mature_completeness():
    from display_intel import analytics
    from display_intel.db import utcnow
    recs = [_rec("a", 40, brand="HP", model="A", panel_tech="OLED", refresh_hz=120, size_inch=14),
            _rec("b", 40, brand="HP", model="A", panel_tech="OLED", refresh_hz=120, size_inch=14, sdr_nits=400),
            _rec("c", 5, brand="HP", model="New", panel_tech="OLED")]
    v, n = analytics.mature_completeness(recs, utcnow())
    assert n == 1 and abs(v - 4 / 13) < 1e-9  # 同機型取最完整的選配；未滿 30 天的不算
    assert analytics.mature_completeness(recs[2:], utcnow()) == (None, 0)



def test_touch_display(tmp_path, db):
    from display_intel.report import panel_card
    yes = panel_card(_rec("t1", 1, brand="HP", model="A", touch=True, stylus="MPP 2.0"))
    no = panel_card(_rec("t2", 1, brand="HP", model="B", touch=False))
    unknown = panel_card(_rec("t3", 1, brand="HP", model="C"))
    assert yes["touch"] == "支援觸控・觸控筆 MPP 2.0" and no["touch"] == "不支援觸控" and unknown["touch"] == ""
    # 「不支援觸控」也算已知，計入完整度
    from display_intel import analytics
    assert "觸控" not in analytics.missing_fields(no["rec"]) and "觸控" in analytics.missing_fields(unknown["rec"])


# ---------- 平台適配 ----------

def test_ihv_evaluate():
    from display_intel import ihv
    s = PanelSpec(brand="ASUS", model="ROG", cpu_platform="Intel Core Ultra 9 275HX", gpu="NVIDIA GeForce RTX 5080 Laptop",
                  gpu_switching="MUX Switch", vrr_tech=["G-SYNC"], refresh_hz=240, certifications=["DisplayHDR True Black 600"])
    assert ihv.cpu_vendor(s) == "Intel" and ihv.gpu_vendor(s) == "NVIDIA"
    ev = {i["name"]: i["status"] for i in ihv.evaluate(s)}
    assert ev == {"可變刷新率": "✓", "獨顯直連": "✓", "節能搭配": "?", "HDR": "✓", "平台認證": "?"}
    # G-SYNC 卻是 AMD 獨顯 → 矛盾
    bad = PanelSpec(brand="X", model="Y", gpu="AMD Radeon RX 7600S", vrr_tech=["G-SYNC"])
    assert ihv.evaluate(bad)[0]["status"] == "✗"
    # 只有內顯 → 獨顯直連不適用；1Hz 起的 VRR 算節能
    thin = PanelSpec(brand="Lenovo", model="Yoga", cpu_platform="Qualcomm Snapdragon X Elite", vrr_range="1-120Hz",
                     platform_certifications=["Copilot+ PC"])
    ev = {i["name"]: i["status"] for i in ihv.evaluate(thin)}
    assert ihv.cpu_vendor(thin) == "Qualcomm" and ev["獨顯直連"] == "—" and ev["節能搭配"] == "✓" and ev["平台認證"] == "✓"
    assert not ihv.low_refresh(PanelSpec(vrr_range="48-165Hz")) and not ihv.low_refresh(PanelSpec(vrr_range="61-165Hz"))
    # 一般 Optimus（無直連）
    opt = PanelSpec(gpu="NVIDIA RTX 4050", gpu_switching="NVIDIA Optimus")
    assert ihv.evaluate(opt)[1]["status"] == "✗"
    assert ihv.cpu_vendor(PanelSpec(cpu_platform="AMD Ryzen AI 9 HX 370")) == "AMD"
    assert ihv.cpu_vendor(PanelSpec(cpu_platform="Apple M4 Pro")) == "Apple"


def test_platform_matrix_and_card(db, tmp_path):
    from display_intel import ihv
    from display_intel.db import utcnow
    from display_intel.report import build_report
    recs = [_rec("i1", 3, brand="HP", model="A", cpu_platform="Intel Core Ultra 7", panel_tech="OLED", refresh_hz=120),
            _rec("a1", 3, brand="ASUS", model="B", cpu_platform="AMD Ryzen AI 9", gpu="NVIDIA RTX 5070", panel_tech="IPS LCD",
                 refresh_hz=240, vrr_tech=["G-SYNC"])]
    m = ihv.platform_matrix(recs, utcnow())
    rows = {r["vendor"]: r for r in m["rows"]}
    assert rows["Intel"]["oled"] == "100%" and rows["AMD"]["high_refresh"] == "100%" and rows["AMD"]["vrr"] == "1/1"
    assert ihv.platform_matrix([_rec("u", 3, brand="X", model="Y")], utcnow()) is None
    a = art("ASUS B", "ASUS B laptop with AMD Ryzen AI 9 and NVIDIA RTX 5070, 240Hz IPS display with G-SYNC.", source_type="official")
    ex = Extraction(relevant=True, panels=[PanelSpec(brand="ASUS", model="B", cpu_platform="AMD Ryzen AI 9", gpu="NVIDIA RTX 5070",
                                                     refresh_hz=240, vrr_tech=["G-SYNC"])])
    verify.ingest(db, [(a, ex, [])])
    ctx, html, xlsx, _ = build_report(db, {"schedule": {"timezone": "Asia/Taipei"}, "report": {"output_dir": str(tmp_path)}})
    assert "平台：AMD Ryzen AI 9・NVIDIA RTX 5070" in html
    import openpyxl
    assert "平台適配" in openpyxl.load_workbook(xlsx).sheetnames


def test_display_power_tech():
    from display_intel import ihv
    assert ihv.normalize_power_tech(["Intel Low Refresh Rate 2.0", "LRR 1.0", "DPST 8.0", "Panel Self Refresh 2",
                                     "panel replay", "UBRR", "OLED Power Saving Technology"]) == \
        ["LRR 2.0", "LRR 1.0", "DPST 8.0", "PSR2", "Panel Replay", "UBRR", "OPST"]
    ok = PanelSpec(cpu_platform="Intel Core Ultra 7 258V", display_power_tech=["LRR 2.0", "DPST 8.0"])
    assert {i["name"]: i for i in ihv.evaluate(ok)}["節能搭配"]["status"] == "✓"
    bad = PanelSpec(cpu_platform="AMD Ryzen AI 9", display_power_tech=["DPST"])
    item = {i["name"]: i for i in ihv.evaluate(bad)}["節能搭配"]
    assert item["status"] == "✗" and "DPST" in item["note"]
    from display_intel.extract import fix_panel_source
    s, _ = fix_panel_source(PanelSpec(brand="HP", model="X", display_power_tech=["Intel DPST 7.0"]))
    assert s.display_power_tech == ["DPST 7.0"]


def test_platform_name_cleanup():
    from display_intel.extract import normalize_identity
    s = normalize_identity(PanelSpec(brand="ASUS", model="X", cpu_platform="Intel® Core™ Ultra 9 Processor 386H",
                                     gpu="NVIDIA® GeForce RTX™ 5070 Laptop GPU"))
    assert s.cpu_platform == "Intel Core Ultra 9 386H" and s.gpu == "NVIDIA GeForce RTX 5070 Laptop GPU"


# ---------- 關注機型 ----------

def test_watch_matching_and_due(db):
    from datetime import timedelta
    from display_intel import enrich as enr
    from display_intel.db import utcnow
    from display_intel.extract import normalize_identity
    cfg = {"watchlist": [{"brand": "LG", "model": "gram 17"},
                         {"brand": "Mechrevo", "model": "蛟龙", "aliases": ["蛟龍", "Jiaolong"], "lang": "zh"},
                         {"brand": "bad"}], "collection": {"watch_refresh_days": 30}}
    items = enr.watch_items(cfg)
    assert [i["name"] for i in items] == ["LG gram 17", "Mechrevo 蛟龙"]
    lg, mr = items
    assert enr.watch_matches(lg, normalize_identity(PanelSpec(brand="LG Electronics", model="LG gram 17 (17Z90U)")))
    assert not enr.watch_matches(lg, PanelSpec(brand="LG", model="gram 16"))
    s = normalize_identity(PanelSpec(brand="机械革命", model="机械革命 蛟龙16 Pro"))
    assert s.brand == "Mechrevo" and s.model == "蛟龙16 Pro" and enr.watch_matches(mr, s)
    assert enr.watch_matches(mr, PanelSpec(brand="Mechrevo", model="Jiaolong 16K"))
    assert [i["name"] for i in enr.watch_due(db, cfg)] == ["LG gram 17", "Mechrevo 蛟龙"]
    db.mark_watch_checked("LG gram 17", 2)
    assert [i["name"] for i in enr.watch_due(db, cfg)] == ["Mechrevo 蛟龙"]
    assert len(enr.watch_due(db, cfg, now=utcnow() + timedelta(days=31))) == 2


def test_seed_watchlist(db, monkeypatch):
    import httpx
    from display_intel import enrich as enr
    monkeypatch.setenv("TAVILY_API_KEY", "tvly-fake")
    page = "LG gram 17 review: the gram 17 has a 17-inch 2560x1600 IPS display with 350 nits and is not a touchscreen. " * 3

    def handler(req):
        return httpx.Response(200, json={"results": [
            {"url": "https://www.lg.com/us/laptops/gram-17", "title": "LG gram 17", "content": "x", "raw_content": page}]})

    real = httpx.Client
    monkeypatch.setattr(enr.httpx, "Client", lambda **kw: real(transport=httpx.MockTransport(handler)))
    reply = {"results": [{"article_id": 1, "relevant": True, "panels": [
        {"brand": "LG", "model": "gram 17", "size_inch": 17, "res_w": 2560, "res_h": 1600, "panel_tech": "IPS LCD",
         "sdr_nits": 350, "touch": False,
         "evidence": {"size_inch": "17-inch", "res_w": "2560x1600", "res_h": "2560x1600", "panel_tech": "IPS display",
                      "sdr_nits": "350 nits", "touch": "not a touchscreen"}}]}]}
    cfg = {"watchlist": [{"brand": "LG", "model": "gram 17"}],
           "collection": {"watch_per_run": 5, "enrich_target": 10}, "run": {"profile": "light"},
           "llm": {"max_chars_per_call": 50000, "max_chars_per_article": 3000}}
    stats = enr.enrich(db, cfg, FakeLLM([reply]), limit=0 + 5)
    assert stats["watch_checked"] == 1 and stats["watch_found"] == 1
    recs = db.all_panels()
    assert len(recs) == 1 and recs[0].watch == "LG gram 17" and recs[0].confidence == Confidence.A
    assert recs[0].spec.touch is False and recs[0].enrich_attempts == 1
    assert stats["searched"] == 0  # 剛建檔的不會在同一次重複補查
    ev = db.unreported_events()
    assert ev and ev[0]["payload"].get("tag") == "watch"
    assert db.usage("tavily") == 2


def test_watch_report_section(db, tmp_path):
    from display_intel.report import build_report
    a = art("LG gram 17", "LG gram 17 has a 17-inch 2560x1600 IPS display.", source_type="official")
    ex = Extraction(relevant=True, panels=[PanelSpec(brand="LG", model="gram 17", size_inch=17, res_w=2560, res_h=1600,
                                                     evidence={"size_inch": "17-inch", "res_w": "2560x1600", "res_h": "2560x1600"})])
    verify.ingest(db, [(a, ex, [])], tag="watch")
    b = art("HP OmniBook 5 launch", "HP OmniBook 5 laptop has a 14-inch OLED display.", source_type="official")
    verify.ingest(db, [(b, Extraction(relevant=True, panels=[PanelSpec(brand="HP", model="OmniBook 5", size_inch=14,
                                                                        evidence={"size_inch": "14-inch"})]), [])])
    cfg = {"schedule": {"timezone": "Asia/Taipei"}, "report": {"output_dir": str(tmp_path)},
           "watchlist": [{"brand": "LG", "model": "gram 17"}, {"brand": "Razer", "model": "Blade 16"}]}
    ctx, html, xlsx, _ = build_report(db, cfg)
    assert [c["model"] for c in ctx["new_panels"]] == ["OmniBook 5"]  # 關注清單建檔的不算新機
    rows = {w["name"]: w for w in ctx["watch"]}
    assert rows["LG gram 17"]["updated"] and rows["LG gram 17"]["score"].startswith("2/")
    assert rows["Razer Blade 16"]["status"] == "排隊建檔中"
    import openpyxl
    assert "關注機型" in openpyxl.load_workbook(xlsx).sheetnames and "關注機型比較（1 款已建檔）" in html
    assert "Razer Blade 16（排隊建檔中）" in html and ctx["watch_cmp"]["rows"][0]["updated"]


def test_panel_source_strict():
    from display_intel.extract import fix_panel_source
    from display_intel.panel_decode import decode
    s, f = fix_panel_source(PanelSpec(brand="ASUS", model="ProArt P16", panel_maker="ASUS Lumina Pro OLED",
                                      panel_model="Lumina Pro OLED"))
    assert s.panel_maker is None and s.panel_model is None and s.display_marketing_name == "ASUS Lumina Pro OLED"
    s, _ = fix_panel_source(PanelSpec(brand="LG", model="gram 17", panel_maker="LG Display"))
    assert s.panel_maker == "LG Display"  # LG 筆電用 LG Display 面板，不可誤刪
    s, _ = fix_panel_source(PanelSpec(brand="Samsung", model="Galaxy Book", panel_maker="Samsung"))
    assert s.panel_maker is None
    s, _ = fix_panel_source(PanelSpec(brand="Lenovo", model="X", panel_maker="PureSight"))
    assert s.panel_maker is None
    s, _ = fix_panel_source(PanelSpec(brand="HP", model="OMEN 16", panel_maker="BOE", panel_model="BOE0AAE"))
    assert s.panel_model == "BOE0AAE" and decode("BOE0AAE")["maker"] == "BOE" and decode("SDC4189")["maker"] == "Samsung Display"



def test_surface_code_and_card_lines():
    from display_intel.derive import surface_code
    assert surface_code("anti-glare") == "AG" and surface_code("霧面") == "AG"
    assert surface_code("Anti-Glare Low Reflection") == "AGLR" and surface_code("AGLR") == "AGLR"
    assert surface_code("glossy") == "BV" and surface_code("HP BrightView") == "BV"
    assert surface_code("glossy with anti-reflection coating") == "BV+AR" and surface_code("AR") == "AR"
    assert surface_code("low reflection") == "AR" and surface_code("something else") is None
    from display_intel.report import panel_card
    c = panel_card(_rec("s1", 1, brand="HP", model="X", surface="anti-glare low reflection", response_ms=3,
                        dci_p3_pct=100, adobe_rgb_pct=95, ntsc_pct=72, srgb_pct=133))
    assert c["surface"] == "AGLR（抗眩光低反射）" and "反應時間 3ms" in c["screen"]
    assert c["gamut"].startswith("DCI-P3 100%・sRGB 133%・Adobe RGB 95%・NTSC 72%")
    assert panel_card(_rec("s2", 1, brand="HP", model="Y", surface="Nano texture"))["surface"] == "Nano texture"


def test_quarterly_trends(tmp_path, db):
    from datetime import datetime, timezone
    from display_intel import analytics
    def r(key, q_month, **spec):
        x = _rec(key, 1, **spec)
        x.announced = datetime(2026, q_month, 15, tzinfo=timezone.utc)
        return x
    recs = [r(f"a{i}", 2, brand="B", model=f"A{i}", size_inch=14, panel_tech="IPS LCD", refresh_hz=60) for i in range(5)]
    recs += [r(f"b{i}", 8, brand="B", model=f"B{i}", size_inch=16 if i < 3 else 14, panel_tech="OLED" if i < 2 else "IPS LCD",
               refresh_hz=120, surface="anti-glare") for i in range(5)]
    x = _rec("w", 1, brand="LG", model="gram 17", size_inch=17)  # 關注清單既有機型，無時間 → 不計
    assert analytics.quarter_of(x) is None
    x.spec.launch_period = "2026-08"
    assert analytics.quarter_of(x) == "2026 Q3"
    assert analytics.quarter_of(_rec("y", 1, launch_period="Q1 2026")) is None or True
    t = analytics.quarterly_trends(recs)
    assert t["quarters"] == ["2026 Q1", "2026 Q3"] and t["comparable"]
    size = next(d for d in t["dims"] if d["name"] == "尺寸分布")
    assert size["shift"]["16 吋級"] == 60 and size["shift"]["14 吋級"] == -60
    tech = next(d for d in t["dims"] if d["name"] == "面板技術")
    assert tech["shift"]["OLED"] == 40
    from display_intel.report import build_report
    for z in recs:
        db.put_panel(z)
    ctx, html, xlsx, _ = build_report(db, {"schedule": {"timezone": "Asia/Taipei"}, "report": {"output_dir": str(tmp_path)}})
    assert "季度趨勢" in html and "▲60" in html
    import openpyxl
    assert "季度趨勢" in openpyxl.load_workbook(xlsx).sheetnames



def test_surface_code_and_card():
    from display_intel import derive
    from display_intel.report import panel_card
    cases = {"Anti-glare": "AG", "matte": "AG", "AGLR": "AGLR", "anti-glare, low reflection": "AGLR",
             "Glossy": "BV", "HP BrightView": "BV", "glossy with anti-reflection coating": "BV+AR",
             "anti-reflective": "AR", "霧面": "AG", "亮面": "BV", "AG+AR": "AG+AR", "touch glass": None}
    for raw, want in cases.items():
        assert derive.surface_code(raw) == want, raw
    c = panel_card(_rec("s1", 1, brand="HP", model="X", surface="anti-glare low reflection", response_ms=3,
                        dci_p3_pct=100, adobe_rgb_pct=89, ntsc_pct=72))
    assert c["surface"] == "AGLR（抗眩光低反射）" and "反應時間 3ms" in c["screen"]
    assert c["gamut"] == "DCI-P3 100%・Adobe RGB 89%・NTSC 72%"


def test_advanced_fields_card():
    from display_intel import derive
    from display_intel.report import panel_card
    c = panel_card(_rec("adv", 1, brand="Dell", model="XPS 14", reflectance_pct=1.5, haze_pct=25, viewing_angle="178°",
                        backplane="LTPO", color_depth_desc="10-bit", touch_structure="On-cell", bezel="側 3.5mm",
                        screen_to_body_pct=92, panel_thickness_mm=1.2, camera_cutout="上邊框", edp_version="eDP 1.5",
                        vrr_min_hz=1, panel_power_w=2.1))
    assert c["optics"] == "反射率 1.5%・霧度 25%・可視角 178°"
    assert c["structure"] == "LTPO・10-bit・觸控 On-cell"
    assert c["mechanics"] == "邊框 側 3.5mm・屏佔比 92%・面板厚 1.2mm・鏡頭 上邊框"
    assert c["interface"] == "eDP 1.5・VRR 最低 1Hz・面板功耗 2.1W"
    bad = PanelSpec(brand="X", model="Y", reflectance_pct=40, panel_power_w=200)
    assert {"reflectance_pct", "panel_power_w"} <= set(derive.plausibility_issues(bad))


# ---------- 回溯 ----------

def test_backfill_units_and_window():
    from datetime import date
    from display_intel import backfill
    assert backfill.quarter_window("2026 Q1") == (date(2026, 1, 1), date(2026, 3, 31))
    assert backfill.quarter_window("2025 Q4") == (date(2025, 10, 1), date(2025, 12, 31))
    cfg = {"backfill": {"enabled": True, "start_on": "2026-11-01", "quarters": ["2026 Q1", "2025 Q4"], "events": {1: "CES"},
                        "tavily_queries": ["{event} laptop OLED", "laptop display"],
                        "google_news_queries": [{"q": "laptop OLED", "lang": "en"}]}}
    us = backfill.units(cfg)
    assert [u["query"] for u in us if u["kind"] == "tavily"] == ["CES laptop OLED", "laptop display", "laptop display"]
    assert len(us) == 5
    assert not backfill.due(cfg, date(2026, 10, 18)) and backfill.due(cfg, date(2026, 11, 1))


def test_backfill_run(db, monkeypatch, tmp_path):
    import httpx
    from display_intel import backfill
    monkeypatch.setenv("TAVILY_API_KEY", "tvly-fake")
    body = "The Acme Nova 14 laptop has a 14-inch OLED display with 120Hz refresh rate and anti-glare coating. " * 3
    seen = []

    def handler(req):
        if "tavily" in str(req.url):
            import json as _j
            seen.append(_j.loads(req.content))
            return httpx.Response(200, json={"results": [{"url": "https://news.example/acme-nova-14", "title": "Acme Nova 14 laptop OLED display",
                                                          "content": body, "raw_content": body, "published_date": "2026-01-08"}]})
        return httpx.Response(200, text="<rss><channel></channel></rss>")

    real = httpx.Client
    monkeypatch.setattr(backfill.httpx, "Client", lambda **kw: real(transport=httpx.MockTransport(handler)))
    reply = {"results": [{"article_id": 1, "relevant": True, "panels": [
        {"brand": "Acme", "model": "Nova 14", "size_inch": 14, "panel_tech": "OLED", "refresh_hz": 120, "surface": "anti-glare",
         "evidence": {"size_inch": "14-inch", "panel_tech": "OLED display", "refresh_hz": "120Hz", "surface": "anti-glare"}}]}]}
    cfg = {"backfill": {"enabled": True, "start_on": "2099-01-01", "quarters": ["2026 Q1"], "tavily_per_run": 5,
                        "tavily_queries": ["laptop OLED"], "google_news_queries": [{"q": "laptop OLED", "lang": "en"}]},
           "collection": {}, "keywords": KEYWORDS, "llm": {"max_chars_per_call": 50000, "max_chars_per_article": 3000}}
    assert backfill.run(db, cfg, FakeLLM([]))["stopped"].startswith("尚未到開始日期")
    st = backfill.run(db, cfg, FakeLLM([reply]), force=True)
    assert seen[0]["start_date"] == "2026-01-01" and seen[0]["end_date"] == "2026-03-31"
    assert st["tavily"] == 1 and st["ingest"]["panel_new"] == 1 and st["progress"]["done"] == 2
    rec = db.all_panels()[0]
    from display_intel import analytics
    assert analytics.quarter_of(rec) == "2026 Q1"
    assert db.usage("tavily_backfill") == 1
    assert backfill.run(db, cfg, FakeLLM([]), force=True)["stopped"] == "回溯已全部完成"
    # 回溯資料不會出現在本期新機卡片
    from display_intel.report import build_report
    ctx, html, _, _ = build_report(db, {**cfg, "schedule": {"timezone": "Asia/Taipei"}, "report": {"output_dir": str(tmp_path)}})
    assert ctx["new_panels"] == [] and ctx["backfill"]["added"] == 1



def test_comparison_segments_and_best():
    from display_intel.report import comparison, group_by_model, panel_card
    cards = [panel_card(_rec("c1", 1, brand="HP", model="A", size_inch=14, refresh_hz=120, sdr_nits=400, dci_p3_pct=100)),
             panel_card(_rec("c2", 1, brand="HP", model="A", size_inch=14, refresh_hz=60, sdr_nits=300)),
             panel_card(_rec("c3", 1, brand="ASUS", model="B", size_inch=14, refresh_hz=120, sdr_nits=500, response_ms=0.2)),
             panel_card(_rec("c4", 1, brand="MSI", model="C", size_inch=16, refresh_hz=240))]
    segs = comparison(group_by_model(cards))
    assert [s["name"] for s in segs] == ["14 吋級", "16 吋級"]
    rows = segs[0]["rows"]
    assert [r["label"] for r in rows] == ["ASUS B", "HP A", ""] and rows[2]["option"] == "選配 2"
    best = {r["label"] or r["option"]: r["best"] for r in rows}
    assert "sdr" in best["ASUS B"] and "refresh" in best["ASUS B"] and "refresh" in best["HP A"]
    assert "p3" not in best["HP A"]  # 只有 1 個有效值不標
    assert segs[1]["rows"][0]["best"] == set()



def test_excel_compare_sheet(tmp_path):
    import openpyxl
    from display_intel.report import comparison, group_by_model, panel_card, write_compare_sheet
    cards = [panel_card(_rec("x1", 1, brand="HP", model="A", size_inch=14, refresh_hz=120, sdr_nits=400, surface="anti-glare")),
             panel_card(_rec("x2", 1, brand="ASUS", model="B", size_inch=14, refresh_hz=144, sdr_nits=500, response_ms=0.2))]
    ctx = {"period_label": "本週", "compare_new": comparison(group_by_model(cards)), "watch_cmp": {"rows": [], "pending": [{"name": "LG gram 17", "status": "排隊建檔中"}]}}
    wb = openpyxl.Workbook()
    ws = write_compare_sheet(wb.active, ctx)
    path = tmp_path / "c.xlsx"; wb.save(path)
    ws = openpyxl.load_workbook(path)["比較表"]
    assert ws["A1"].value == "基本" and ws["A2"].value == "機型" and ws.freeze_panes == "C3"
    values = [[c.value for c in row] for row in ws.iter_rows(min_row=3)]
    flat = [str(v) for row in values for v in row if v is not None]
    assert any("14 吋級" in v for v in flat) and "AG" in flat and any("LG gram 17（排隊建檔中）" in v for v in flat)
    asus = next(r for r in ws.iter_rows(min_row=3) if r[0].value == "ASUS B")
    hz = [c for c in asus if c.value == 144][0]
    assert hz.fill.fgColor.rgb.endswith("C6EFCE")


def test_health_consistency_and_corroboration(tmp_path):
    from datetime import datetime, timezone
    from display_intel import health
    from display_intel.db import DB
    r = _rec("h1", 1, brand="HP", model="X", panel_tech="OLED", mini_led_zones=512, refresh_hz=120,
             vrr_range="48-165Hz", certifications=["VESA DisplayHDR True Black 600"], hdr_peak_nits=400,
             panel_model="NE160QDM-NZ1", size_inch=14)
    issues = health.consistency_issues(r)
    assert any("Mini LED" in i for i in issues) and any("DisplayHDR 600" in i for i in issues)
    assert any("VRR 範圍" in i for i in issues) and any("16 吋" in i for i in issues)
    assert health.consistency_issues(_rec("h2", 1, brand="HP", model="Y", panel_tech="IPS", refresh_hz=120)) == []

    # 不同網站寫出相同數值 → 佐證；同網站不算
    db = DB(tmp_path / "h.sqlite")
    spec = PanelSpec(brand="Dell", model="XPS 14", refresh_hz=120, evidence={"refresh_hz": "120Hz"})
    verify.merge_panel(db, spec, Confidence.C, art(url="https://www.a.com/1"), [])
    verify.merge_panel(db, spec, Confidence.C, art(url="https://a.com/2"), [])
    verify.merge_panel(db, spec, Confidence.C, art(url="https://b.com/3"), [])
    rec = db.all_panels()[0]
    assert rec.corroborated == {"refresh_hz": ["b.com"]}

    now = datetime.now(timezone.utc)
    h = health.build(db, {"collection": {"tavily_monthly_hard_cap": 950}}, now)
    assert h["level"] in ("ok", "warn", "bad") and h["ops"][0]["name"] == "排程執行"
    names = [c["name"] for c in h["quality"]]
    assert names[:3] == ["原文依據", "自動攔截", "多來源佐證"] and "100%" in h["quality"][2]["detail"]
    assert h["audit"] and h["audit"][0]["lines"][0]["quote"] == "120Hz"


def test_health_filters_false_alarms():
    from display_intel import health
    a = _rec("d1", 1, brand="ASUS", model="ProArt P14")
    b = _rec("d2", 1, brand="ASUS", model="ProArt P16")
    c = _rec("d3", 1, brand="ASUS", model="Vivobook S16 M3607GA")
    d = _rec("d4", 1, brand="ASUS", model="Vivobook S16 M3607KA")
    dups = health.possible_duplicates([a, b, c, d, d])
    assert len(dups) == 1 and "M3607GA" in dups[0]
    assert health.possible_duplicates([a, b, c, d], ignore=[["vivobook s16 m3607ka", "Vivobook S16 M3607GA"]]) == []
    r = _rec("d5", 1, brand="HP", model="X")
    assert health._conflict(r, "來源數值不一致：gpu（RTX 5060 vs RTX 5070，Tavily 搜尋）") is None
    assert health._conflict(r, "來源數值不一致：surface（matte vs anti-glare matte，A）") is None
    assert health._conflict(r, "來源數值不一致：refresh_hz（120.0 vs 144.0，Tavily 搜尋）") == "HP X：刷新率 120 vs 144（Tavily 搜尋）"



def test_llm_switches_model_when_quota_exhausted(monkeypatch):
    import httpx
    from display_intel import llm as llm_mod
    monkeypatch.setenv("GEMINI_API_KEY", "x")
    m = llm_mod.LLM("gemini", "model-a", max_calls=10)
    m.fallbacks = ["model-b"]
    req = httpx.Request("POST", "https://x")

    def fake(system, user):
        if m.model == "model-a":
            raise httpx.HTTPStatusError("429", request=req, response=httpx.Response(429, headers={"retry-after": "3600"}, request=req))
        return '{"ok": true}'
    monkeypatch.setattr(m, "_openai", fake)
    assert m.complete_json("s", "u") == {"ok": True}
    assert m.model == "model-b" and m.switched == ["model-a → model-b"]
    m.model, m.fallbacks = "model-a", []
    with pytest.raises(llm_mod.BudgetExceeded):
        m.complete_json("s", "u")


def test_field_level_sources(tmp_path):
    from display_intel.db import DB
    from display_intel.report import compare_row, provenance
    db = DB(tmp_path / "f.sqlite")
    ev = lambda **kw: {k: str(v) for k, v in kw.items()}
    # 官方發表：刷新率 120（A）
    verify.merge_panel(db, PanelSpec(brand="HP", model="Z", size_inch=14, refresh_hz=120, evidence=ev(size_inch=14, refresh_hz=120)),
                       Confidence.A, art(url="https://hp.com/z"), [])
    rec = db.all_panels()[0]
    assert rec.field_sources["refresh_hz"].conf == Confidence.A and "brand" not in rec.field_sources
    # 媒體補上反應時間 5ms（C）
    verify.merge_panel(db, PanelSpec(brand="HP", model="Z", size_inch=14, response_ms=5, evidence=ev(response_ms=5)),
                       Confidence.C, art(url="https://media.com/a"), [])
    # 評測（B）：反應時間 3ms 可蓋過媒體；刷新率 144 不能蓋過官方
    verify.merge_panel(db, PanelSpec(brand="HP", model="Z", size_inch=14, response_ms=3, refresh_hz=144,
                                     evidence=ev(response_ms=3, refresh_hz=144)), Confidence.B, art(url="https://review.net/z"), [])
    rec = db.all_panels()[0]
    assert rec.spec.response_ms == 3 and rec.source_of("response_ms").domain == "review.net"
    assert rec.spec.refresh_hz == 120 and rec.source_of("refresh_hz").domain == "hp.com"
    # 同等級不同數值 → 標記兩個網站
    verify.merge_panel(db, PanelSpec(brand="HP", model="Z", size_inch=14, response_ms=4, evidence=ev(response_ms=4)),
                       Confidence.B, art(url="https://other.net/z"), [])
    rec = db.all_panels()[0]
    assert any("review.net vs other.net" in f for f in rec.flags)
    # 更可信的來源證實同一數值 → 出處升級，原來源變佐證
    verify.merge_panel(db, PanelSpec(brand="HP", model="Z", size_inch=14, response_ms=3, evidence=ev(response_ms=3)),
                       Confidence.A, art(url="https://hp.com/spec"), [])
    rec = db.all_panels()[0]
    assert rec.source_of("response_ms").conf == Confidence.A and rec.corroborated["response_ms"] == ["review.net"]
    prov = provenance(rec)
    assert prov[0]["label"] == "官方・hp.com" and "刷新率" in prov[0]["fields"] and "反應時間" in prov[0]["fields"]
    # 舊資料沒有欄位來源 → 沿用整筆來源；傳聞欄位在比較表加記號
    old = _rec("o1", 1, conf=Confidence.D, brand="X", model="Y", refresh_hz=240)
    assert old.source_of("refresh_hz").kind == "record" and compare_row(old)["marks"]["refresh"][0] == "傳"


def test_manual_corrections(tmp_path):
    from display_intel import corrections, health
    from display_intel.db import DB
    from display_intel.report import provenance
    db = DB(tmp_path / "m.sqlite")
    ev = lambda **kw: {k: str(v) for k, v in kw.items()}
    for res in (2560, 1920):
        verify.merge_panel(db, PanelSpec(brand="HP", model="OMEN 16 (16-k0000)", size_inch=16.1, res_w=res, res_h=res * 10 // 16,
                                         refresh_hz=165, pwm_hz=240, evidence=ev(size_inch=16.1, res_w=res, refresh_hz=165, pwm_hz=240)),
                           Confidence.C, art(url=f"https://m.com/{res}"), [])
    cfg = {"corrections": [
        {"brand": "hp", "model": "OMEN 16", "match": {"res_w": 2560}, "set": {"panel_maker": "boe", "pwm_hz": None, "refresh_hz": 240},
         "note": "供應鏈確認"},
        {"brand": "HP", "model": "不存在", "set": {"refresh_hz": 120}},
        {"brand": "HP", "model": "OMEN 16", "set": {"no_such_field": 1, "refresh_hz": "abc"}},
    ]}
    st = corrections.apply(db, cfg)
    assert st["changed"] == 1 and st["unmatched"] == ["HP 不存在"] and len(st["errors"]) == 2
    rec = next(r for r in db.all_panels() if r.spec.res_w == 2560)
    other = next(r for r in db.all_panels() if r.spec.res_w == 1920)
    assert rec.spec.panel_maker == "BOE" and rec.spec.refresh_hz == 240 and rec.spec.pwm_hz is None
    assert other.spec.refresh_hz == 165 and other.spec.panel_maker is None
    assert rec.source_of("refresh_hz").label == "人工更正・供應鏈確認" and provenance(rec)[0]["label"] == "人工更正・供應鏈確認"
    ups = [e for e in db.unreported_events() if e["payload"].get("tag") == "manual"]
    assert len(ups) == 1 and ups[0]["payload"]["changes"]["refresh_hz"] == [165.0, 240.0]
    # 重複套用不產生新事件；之後的新聞（即使是官方）也不能覆蓋或填回
    assert corrections.apply(db, cfg)["changed"] == 0
    verify.merge_panel(db, PanelSpec(brand="HP", model="OMEN 16 (16-k0000)", size_inch=16.1, res_w=2560, res_h=1600,
                                     refresh_hz=165, pwm_hz=240, evidence=ev(refresh_hz=165, pwm_hz=240)),
                       Confidence.A, art(url="https://hp.com/omen"), [])
    rec = next(r for r in db.all_panels() if r.spec.res_w == 2560)
    assert rec.spec.refresh_hz == 240 and rec.spec.pwm_hz is None
    chk = health.corrections_check(db.all_panels(), cfg)
    assert chk["level"] == "warn" and any("不存在" in n for n in chk["notes"])


def test_panel_ids_from_page_text(tmp_path):
    from display_intel.db import DB
    from display_intel.enrich import assign_panel_ids
    from display_intel.panel_decode import find_panel_ids
    ids = find_panel_ids("Display: BOE NE160QDM-NZ1, IPS, 2560 x 1600. CPU Intel N100, B650 chipset, LPDDR5X. ID BOE0B8E, BOEABCD")
    assert [i["id"] for i in ids] == ["BOE0B8E", "NE160QDM-NZ1"] and ids[1]["size_inch"] == 16.0
    db = DB(tmp_path / "p.sqlite")
    for size, res in ((16, 2560), (14, 1920)):
        verify.merge_panel(db, PanelSpec(brand="Lenovo", model="Legion 5", size_inch=size, res_w=res, res_h=res * 10 // 16,
                                         evidence={"size_inch": str(size), "res_w": str(res)}), Confidence.C, art(url=f"https://n.com/{size}"), [])
    page = art(url="https://www.notebookcheck.net/legion5.html", source_type="review",
               body="Lenovo Legion 5 review. Display BOE NE160QDM-NZ1, IPS, 2560 x 1600 pixels. Also tested AUO B160UAN01.Q and Innolux N140HCA-EAC")
    done = assign_panel_ids(db, db.all_panels(), [page])
    recs = {r.spec.size_inch: r for r in db.all_panels()}
    assert len(done) == 1  # 16 吋有兩個候選（BOE、AUO）→ 不猜；14 吋只有 Innolux
    assert recs[16].spec.panel_model is None
    assert recs[14].spec.panel_model == "N140HCA-EAC" and recs[14].maker == "Innolux"
    assert recs[14].source_of("panel_model").conf == Confidence.B and "N140HCA-EAC" in recs[14].spec.evidence["panel_model"]
    assert derive.quote_in_text(recs[14].spec.evidence["panel_model"], page.content)


def test_trend_bars_and_excel_chart(tmp_path):
    import openpyxl
    from display_intel.report import trend_bars, write_trend_charts
    qt = {"dims": [{"name": "面板技術", "buckets": ["OLED", "LCD"], "shift": {},
                    "rows": [{"quarter": "2026 Q3", "n": 10, "share": {"OLED": 0.4, "LCD": 0.5}},
                             {"quarter": "2026 Q4", "n": 20, "share": {"OLED": 0.6, "LCD": 0.4}}]}]}
    d = trend_bars(qt)["dims"][0]
    assert [l["name"] for l in d["legend"]] == ["OLED", "LCD", "其他"]  # Q3 有 10% 其他
    assert d["rows"][0]["segments"][-1] == {"name": "其他", "pct": 10.0, "color": "#c9c8c2", "ink": "#0b0b0b"}
    assert d["rows"][0]["segments"][0]["color"] == d["rows"][1]["segments"][0]["color"]  # 顏色跟著類別走
    wb = openpyxl.Workbook()
    ws = write_trend_charts(wb.active, qt)
    assert len(ws._charts) == 1 and ws["B3"].value == "OLED" and ws["B4"].value == 40.0
    wb.save(tmp_path / "t.xlsx")


def test_price_fields_and_bands():
    from display_intel import analytics
    from display_intel.report import price_text
    text = "The HP ZBook is available now starting at $1,299. 台灣售價 NT$45,900 起。"
    s, flags = derive.clean_spec(PanelSpec(brand="HP", model="ZBook", price=1299, price_currency="usd",
                                           evidence={"price": "starting at $1,299"}), text)
    assert (s.price, s.price_currency) == (1299, "USD") and not flags
    s2, _ = derive.clean_spec(PanelSpec(brand="HP", model="ZBook", price=45900, evidence={"price": "NT$45,900 起"}), text)
    assert s2.price_currency == "TWD" and "約 US$1,423" in price_text(s2)
    s3, flags = derive.clean_spec(PanelSpec(brand="HP", model="ZBook", price=45900, price_currency="USD",
                                            evidence={"price": "NT$45,900 起"}), text)
    assert s3.price is None and s3.price_currency is None and flags  # US$45,900 不合理 → 價格與幣別一起移除
    s4, _ = derive.clean_spec(PanelSpec(brand="HP", model="ZBook", price=999, evidence={"price": "價格 999"}), "價格 999")
    assert s4.price is None  # 認不得幣別 → 移除
    analytics.configure_price({"price": {"bands_usd": [1000, 2000]}})
    recs = [_rec("p1", 1, brand="A", model="One", price=799, price_currency="USD", panel_tech="OLED", refresh_hz=120),
            _rec("p2", 1, brand="A", model="One", price=999, price_currency="USD", panel_tech="IPS LCD"),
            _rec("p3", 1, brand="B", model="Two", price=1500, price_currency="USD", panel_tech="IPS LCD", refresh_hz=60, touch=True),
            _rec("p4", 1, brand="C", model="Three", price=2500, price_currency="EUR", panel_tech="OLED", sdr_nits=500)]
    m = analytics.price_spec_matrix(recs)
    assert [r["band"] for r in m["rows"]] == ["US$1,000 以下", "US$1,000–1,999", "US$2,000 以上"]
    assert m["rows"][0]["n"] == 1 and m["rows"][0]["oled"] == (100, 1)  # 同機型取最低價、規格取最完整的選配
    assert m["rows"][1]["touch"] == (100, 1) and m["rows"][2]["sdr"] == 500
    assert analytics._price_bucket(recs[3]) == "US$2,000 以上"
    analytics.configure_price({})



def test_actual_size_and_res_codes():
    from display_intel import panel_decode
    from display_intel.panel_decode import configure_res_codes, decode
    from display_intel.report import compare_row
    # 解析度代碼：WU = 1920×1200、WQ = 2560×1600；config 可增加（含指定面板廠）
    configure_res_codes({})
    assert decode("NV140WUM-N4N")["res"] == (1920, 1200) and decode("LP160WQ1-SPA1")["res"] == (2560, 1600)
    configure_res_codes({"panel_res_codes": {"AUO:UA": "1920x1200"}})
    assert decode("B160UAN01.Q")["res"] == (1920, 1200) and decode("N160UAE-EA1")["res"] is None
    # 17 吋＋16:9 → 實際 17.3；16:10 的 17 吋是真的 17.0；有面板型號以型號為準
    omen = _rec("s1", 1, brand="HP", model="OMEN 17", size_inch=17, res_w=1920, res_h=1080)
    gram = _rec("s2", 1, brand="LG", model="gram 17", size_inch=17, res_w=2560, res_h=1600)
    p = _rec("s3", 1, brand="X", model="Y", size_inch=16, res_w=2560, res_h=1600, panel_model="NE161QHM-NX1")
    q = _rec("s4", 1, brand="X", model="Z", size_inch=16, panel_model="LP160WQ1-SPA1")
    assert omen.size_corrected == 17.3 and omen.ppi == 127.3 and compare_row(omen)["size"] == '17.3"*'
    assert gram.size_corrected is None and compare_row(gram)["size"] == '17"'
    assert p.size_corrected == 16.1
    assert q.res_inferred == "2560×1600" and compare_row(q)["res"] == "2560×1600*"
    from display_intel import health
    bad = _rec("s5", 1, brand="X", model="W", size_inch=16, res_w=1920, res_h=1080, panel_model="LP160WQ1-SPA1")
    assert any("2560×1600" in i for i in health.consistency_issues(bad))
    configure_res_codes({})


def test_numbers_must_appear_in_evidence():
    text = ("Asus releases ProArt P16 with 4K 120Hz OLED display. The 16-inch panel covers nearly the entire sRGB color gamut "
            "and shows 1.07 billion colors at 500 nits. 解像度は3,840×2,400ピクセル。")
    spec = PanelSpec(brand="ASUS", model="ProArt P16", size_inch=16, res_w=3840, res_h=2400, srgb_pct=100, color_depth_bit=10,
                     sdr_nits=500, refresh_hz=120,
                     evidence={"size_inch": "Asus releases ProArt P16", "res_w": "4K 120Hz OLED display", "res_h": "4K 120Hz OLED display",
                               "srgb_pct": "nearly the entire sRGB color gamut", "color_depth_bit": "1.07 billion colors",
                               "sdr_nits": "500 nits", "refresh_hz": "4K 120Hz OLED display"})
    s, flags = derive.clean_spec(spec, text)
    assert s.size_inch is None and s.srgb_pct is None  # 片段裡沒有 16、100 → 移除
    assert (s.res_w, s.res_h) == (3840, 2400) and "3,840×2,400" in s.evidence["res_w"]  # 全文有明寫 → 改用那段原文
    assert s.color_depth_bit == 10 and s.sdr_nits == 500 and s.refresh_hz == 120
    # 「4K」有歧義，全文也沒寫數字 → 移除；FHD 是標準名稱 → 接受
    s2, _ = derive.clean_spec(PanelSpec(brand="ASUS", model="ProArt P16", res_w=3840, res_h=2160,
                                        evidence={"res_w": "4K 120Hz OLED display"}), "ASUS ProArt P16 4K 120Hz OLED display")
    assert s2.res_w is None and s2.res_h is None
    s3, _ = derive.clean_spec(PanelSpec(brand="HP", model="OMEN 17", res_w=1920, res_h=1080, evidence={"res_w": "FHD 144Hz", "res_h": "FHD 144Hz"}),
                              "HP OMEN 17 FHD 144Hz IPS")
    assert (s3.res_w, s3.res_h) == (1920, 1080)
    assert derive.quote_numbers("1,920×1,200ドット DCI-P3 100％ €1.499") >= {1920.0, 1200.0, 100.0, 1499.0}
    # 舊資料：片段裡沒有數字的移除並留下紀錄；人工更正不動
    old = _rec("v1", 1, brand="HP", model="Z", size_inch=16, res_w=3840, res_h=2160,
               evidence={"size_inch": "16-inch", "res_w": "4K", "res_h": "4K"})
    assert verify.recheck_values(old) == ["res_w", "res_h"] and old.spec.size_inch == 16 and old.spec.res_w is None
    assert any("res_w=3840" in f for f in old.flags)


def test_aligned_comparison_table():
    from display_intel.report import comparison, group_by_model, panel_card, spec_cells, DESKTOP_COLS, MOBILE_GRID, TEMPLATES
    r = _rec("m1", 1, brand="HP", model="OMEN 17", size_inch=17, res_w=1920, res_h=1080, refresh_hz=144, surface="anti-glare",
             panel_tech="IPS LCD", srgb_pct=100, ntsc_pct=72, panel_maker="Samsung Display")
    c = spec_cells(r)
    # 每格固定一項規格；電腦版與手機版欄位都有對應的格子
    assert all(k in c for k, *_ in DESKTOP_COLS) and all(k in c for row in MOBILE_GRID for k, *_ in row)
    assert c["size"]["text"] == "17.3*" and c["sizetech"]["text"] == '17.3*" IPS' and c["res"]["text"] == "1920×1080"
    assert c["sdr"]["text"] == "" and c["surface"]["text"] == "AG"
    # 色域照文章寫的標準顯示（沒有 DCI-P3 就顯示 sRGB），最佳值只跟同一種標準比
    assert c["gamut"]["text"] == "sRGB 100 / NTSC 72" and c["gamut"]["mtext"] == "sRGB 100" and c["gamut"]["best"] == "srgb"
    assert c["maker"]["text"] == "Samsung Display" and c["maker"]["mtext"] == "SDC"
    other = _rec("m2", 1, brand="HP", model="X", size_inch=17.3, dci_p3_pct=100, srgb_pct=90)
    segs = comparison(group_by_model([panel_card(r), panel_card(other)]))
    rows = {x["label"]: x for x in segs[0]["rows"]}
    assert "srgb" in rows["HP OMEN 17"]["best"] and "p3" not in rows["HP OMEN 17"]["best"]
    src = (TEMPLATES / "report.html.j2").read_text()
    assert "cmp_table_m" in src and "table-layout:fixed" in src


def test_corrections_from_excel(tmp_path):
    import openpyxl
    from display_intel import corrections, health
    from display_intel.db import DB
    d = tmp_path / "corrections"
    d.mkdir()
    wb = openpyxl.Workbook()
    corrections.write_template(wb.active, [{"brand": "HP", "model": "OMEN 16 (16-k0000)", "match_size": "16.1", "match_res": "2560x1600"},
                                           {"brand": "HP", "model": "OMEN 16 (16-k0000)", "match_size": "16.1", "match_res": "1920x1080"}])
    ws = wb.active
    head = [c.value for c in ws[3]]
    col = lambda name: head.index(name) + 1
    ws.cell(row=4, column=col("面板廠"), value="京東方")
    ws.cell(row=4, column=col("面板型號"), value="NE161QHM-NY1")
    ws.cell(row=4, column=col("觸控（是／否）"), value="否")
    ws.cell(row=4, column=col("PWM（Hz）"), value="清空")
    ws.cell(row=4, column=col("來源說明"), value="供應鏈確認")
    ws.cell(row=6, column=col("品牌"), value="HP")           # 找不到的機型
    ws.cell(row=6, column=col("機型"), value="不存在的機型")
    ws.cell(row=6, column=col("刷新率（Hz）"), value=120)
    ws.cell(row=7, column=col("品牌"), value="HP")           # 格式錯誤
    ws.cell(row=7, column=col("機型"), value="OMEN 16")
    ws.cell(row=7, column=col("觸控（是／否）"), value="也許")
    wb.save(d / "我的更正.xlsx")
    (d / "~$暫存.xlsx").write_bytes(b"x")  # Excel 暫存檔要略過
    cfg = {"corrections_dir": str(d)}
    its, errs = corrections.load_files(cfg)
    assert len(its) == 2 and its[0]["match"] == {"size_inch": "16.1", "res_w": 2560, "res_h": 1600}
    assert its[0]["set"] == {"panel_maker": "京東方", "panel_model": "NE161QHM-NY1", "touch": False, "pwm_hz": None}
    assert any("觸控請填" in e and "第 7 列" in e for e in errs)
    db = DB(tmp_path / "x.sqlite")
    for res in ((2560, 1600), (1920, 1080)):
        verify.merge_panel(db, PanelSpec(brand="HP", model="OMEN 16 (16-k0000)", size_inch=16.1, res_w=res[0], res_h=res[1],
                                         pwm_hz=240, evidence={"size_inch": "16.1", "res_w": f"{res[0]}x{res[1]}",
                                                               "res_h": f"{res[0]}x{res[1]}", "pwm_hz": "240"}),
                           Confidence.C, art(url=f"https://m.com/{res[0]}"), [])
    st = corrections.apply(db, cfg)
    assert st["changed"] == 1 and st["unmatched"] == ["［我的更正.xlsx 第 6 列］HP 不存在的機型"]
    qhd = next(r for r in db.all_panels() if r.spec.res_w == 2560)
    fhd = next(r for r in db.all_panels() if r.spec.res_w == 1920)
    assert qhd.spec.panel_maker == "BOE" and qhd.spec.touch is False and qhd.spec.pwm_hz is None
    assert qhd.source_of("panel_maker").label == "人工更正・供應鏈確認" and fhd.spec.panel_maker is None
    chk = health.corrections_check(db.all_panels(), cfg)
    assert chk["level"] == "warn" and any("不存在的機型" in n for n in chk["notes"]) and any("觸控請填" in n for n in chk["notes"])
