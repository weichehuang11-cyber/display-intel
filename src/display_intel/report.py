"""報告層：把未報告的事件整理成 HTML Email 與 Excel 附件。"""

from __future__ import annotations

import json
import logging
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import analytics, derive, health, ihv
from .db import DB
from .panel_decode import canonical_maker
from .schema import Confidence, PanelRecord, PanelSpec

log = logging.getLogger(__name__)

TEMPLATES = Path(__file__).parent / "templates"

FIELD_LABELS = {
    "brand": "品牌", "model": "機型", "panel_maker": "面板廠", "panel_model": "面板型號", "panel_tech": "面板技術",
    "size_inch": "尺寸", "res_w": "水平解析度", "res_h": "垂直解析度", "surface": "表面處理",
    "sdr_nits": "SDR 亮度", "hdr_peak_nits": "HDR 峰值亮度", "contrast": "對比", "srgb_pct": "sRGB",
    "dci_p3_pct": "DCI-P3", "adobe_rgb_pct": "Adobe RGB", "color_depth_bit": "色深", "delta_e": "ΔE",
    "hdr_formats": "HDR 格式", "certifications": "認證", "refresh_hz": "刷新率", "vrr_range": "VRR 範圍",
    "response_ms": "反應時間", "vrr_tech": "VRR 技術", "backlight": "背光", "mini_led_zones": "Mini LED 分區",
    "dimming": "調光", "pwm_hz": "PWM 頻率", "low_power_tech": "節能技術", "eye_care": "護眼",
    "ntsc_pct": "NTSC", "reflectance_pct": "反射率 %", "haze_pct": "霧度 %", "uniformity_pct": "亮度均勻度 %",
    "viewing_angle": "可視角", "backplane": "背板", "color_depth_desc": "色深（原文）", "touch_structure": "觸控結構",
    "bezel": "邊框", "panel_thickness_mm": "面板厚度 mm", "camera_cutout": "鏡頭位置", "edp_version": "eDP 版本",
    "vrr_min_hz": "VRR 最低 Hz", "panel_power_w": "面板功耗 W", "touch": "觸控", "stylus": "觸控筆", "glass": "玻璃", "screen_to_body_pct": "屏佔比", "form_factor": "形態",
    "cpu_platform": "處理器", "gpu": "顯示卡", "gpu_switching": "顯示切換", "platform_certifications": "平台認證",
    "display_power_tech": "平台顯示節能",
}
VENDOR_ORDER = ["Intel", "NVIDIA", "AMD", "Qualcomm"]
FAB_STATUS = {"planned": "規劃", "construction": "建廠", "equipment": "設備進駐", "trial": "試產",
              "mass_production": "量產", "shutdown": "停產", "other": "動態"}
PLATFORM_CATEGORY = {"sync": "同步技術", "power": "節能技術", "image": "畫質技術", "interface": "介面標準",
                     "certification": "平台認證", "other": "其他"}


def _fmt(v) -> str:
    if v is None or v == []:
        return ""
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, float):
        return f"{v:g}"
    if isinstance(v, list):
        return "、".join(map(str, v))
    return str(v)


def _current_record(db: DB, ev: dict) -> PanelRecord:
    """事件裡存的是當時的快照；報告改用資料庫中最新、已合併的紀錄。"""
    snap = ev["payload"]["record"]
    rec = db.get_panel(snap["key"])
    if rec is None:
        from .verify import find_existing
        rec = find_existing(db, PanelSpec(**snap["spec"]))
    return rec or PanelRecord.model_validate(snap)


def panel_card(rec: PanelRecord, changes: dict | None = None, tags: list[str] | None = None) -> dict:
    s = rec.spec
    if s.size_inch:
        size = f'{s.size_inch:g}"'
    elif rec.size_inferred:
        size = f'{rec.size_inferred:g}"（推算）'
    else:
        size = ""
    screen = "・".join(x for x in [size, s.panel_tech or "", f"{s.res_w}×{s.res_h}" if s.res_w and s.res_h else "",
                                   f"{s.refresh_hz:g}Hz" if s.refresh_hz else "",
                                   f"反應時間 {s.response_ms:g}ms" if s.response_ms else ""] if x)
    code = derive.surface_code(s.surface)
    surface = (f"{code}（{derive.SURFACE_LABELS.get(code, code)}）" if code in derive.SURFACE_LABELS else code) if code else ""
    if s.surface and not code:
        surface = s.surface  # 認不得的寫法照原文顯示
    bright = "・".join(x for x in [f"SDR {s.sdr_nits:g} nits" if s.sdr_nits else "",
                                   f"HDR {s.hdr_peak_nits:g} nits" if s.hdr_peak_nits else "",
                                   f"對比 {s.contrast}" if s.contrast else ""] if x)
    gamut = "・".join(x for x in [f"DCI-P3 {s.dci_p3_pct:g}%" if s.dci_p3_pct else "",
                                  f"sRGB {s.srgb_pct:g}%" if s.srgb_pct else "",
                                  f"Adobe RGB {s.adobe_rgb_pct:g}%" if s.adobe_rgb_pct else "",
                                  f"NTSC {s.ntsc_pct:g}%" if s.ntsc_pct else "",
                                  f"{s.color_depth_bit}-bit" if s.color_depth_bit else "",
                                  f"ΔE {s.delta_e:g}" if s.delta_e else ""] if x)
    maker = rec.maker
    maker_txt = (f"{maker}（由型號推算）" if rec.panel_maker_inferred else maker) if maker else ""
    panel = " ".join(x for x in [maker_txt, s.panel_model or ""] if x)
    extras = "・".join(x for x in [
        f"品牌命名：{s.display_marketing_name}" if s.display_marketing_name else "",
        _fmt(s.vrr_tech), f"Mini LED {s.mini_led_zones} 區" if s.mini_led_zones else "",
        f"{s.dimming} 調光" if s.dimming else "", f"PWM {s.pwm_hz:g}Hz" if s.pwm_hz else "",
        _fmt(s.low_power_tech), _fmt(s.display_power_tech), _fmt(s.certifications), _fmt(s.hdr_formats)] if x)
    def g(*parts):
        return "・".join(x for x in parts if x)
    optics = g(f"反射率 {s.reflectance_pct:g}%" if s.reflectance_pct else "", f"霧度 {s.haze_pct:g}%" if s.haze_pct else "",
               f"均勻度 {s.uniformity_pct:g}%" if s.uniformity_pct else "", f"可視角 {s.viewing_angle}" if s.viewing_angle else "")
    structure = g(s.backplane or "", s.color_depth_desc or "", f"觸控 {s.touch_structure}" if s.touch_structure else "",
                  f"蓋板 {s.glass}" if s.glass else "")
    mechanics = g(f"邊框 {s.bezel}" if s.bezel else "", f"屏佔比 {s.screen_to_body_pct:g}%" if s.screen_to_body_pct else "",
                  f"面板厚 {s.panel_thickness_mm:g}mm" if s.panel_thickness_mm else "", f"鏡頭 {s.camera_cutout}" if s.camera_cutout else "")
    interface = g(s.edp_version or "", f"VRR 最低 {s.vrr_min_hz:g}Hz" if s.vrr_min_hz else "",
                  f"面板功耗 {s.panel_power_w:g}W" if s.panel_power_w else "")
    if s.touch is True:
        touch = "支援觸控" + (f"・觸控筆 {s.stylus}" if s.stylus else "")
    elif s.touch is False:
        touch = "不支援觸控"
    else:
        touch = f"觸控筆 {s.stylus}" if s.stylus else ""
    derived = "・".join(x for x in [f"{rec.ppi:g} PPI" if rec.ppi else "", rec.aspect_ratio or "",
                                    f"像素密度{rec.density_class}" if rec.density_class else "",
                                    ("面板原廠規格：" + "、".join(FIELD_LABELS.get(f, f) for f in rec.panel_spec_fields))
                                    if rec.panel_spec_fields else ""] if x)
    change_text = "；".join(f"{FIELD_LABELS.get(k, k)}：{_fmt(o) or '—'} → {_fmt(n)}" for k, (o, n) in (changes or {}).items())
    known, total = analytics.completeness(rec)
    return {
        "score": f"{known}/{total}", "score_pct": round(known / total * 100), "missing": "、".join(analytics.missing_fields(rec)),
        "key": rec.key, "rec": rec, "brand": s.brand or "", "model": s.model or "",
        "screen": screen, "bright": bright, "gamut": gamut, "surface": surface, "touch": touch,
        "optics": optics, "structure": structure, "mechanics": mechanics, "interface": interface,
        "platform": ihv.platform_text(s), "ihv": ihv.evaluate(s), "panel": panel or "面板來源待確認", "panel_known": bool(panel),
        "extras": extras, "derived": derived, "tags": tags or [],
        "confidence": rec.confidence.label, "conf_code": rec.confidence.value,
        "source_name": rec.source_name, "source_url": rec.source_url, "changes": change_text,
        "spec": s.model_dump(), "record": rec.model_dump(mode="json"),
        # 舊欄位（AI 摘要與 Excel 使用）
        "size": size, "tech": s.panel_tech or "", "res": f"{s.res_w}×{s.res_h}" if s.res_w and s.res_h else "",
        "refresh": f"{s.refresh_hz:g}Hz" if s.refresh_hz else "", "panel_src": panel or "待確認",
    }


def group_by_model(cards: list[dict]) -> list[dict]:
    """同一款機型的不同螢幕選配合併成一張卡片（例如 OMEN 15 的 FHD 144Hz 與 QHD 165Hz）。"""
    from .verify import _norm, model_norm
    groups: dict[tuple, list[dict]] = {}
    for c in cards:
        k = (_norm(c["rec"].spec.brand), model_norm(c["rec"].spec))
        groups.setdefault(k, []).append(c)
    out = []
    for opts in groups.values():
        opts.sort(key=lambda c: (-c["rec"].confidence.rank, -c["score_pct"]))
        best = max(opts, key=lambda c: c["score_pct"])
        head = dict(opts[0])
        head["options"] = opts
        head["tags"] = list(dict.fromkeys(t for c in opts for t in c["tags"]))
        head["score"], head["score_pct"] = best["score"], best["score_pct"]
        head["panel_known"] = any(c["panel_known"] for c in opts)
        head["sources"] = list({c["source_url"]: c for c in opts if c["source_url"]}.values())[:3]
        out.append(head)
    return sorted(out, key=lambda c: (c["brand"].lower(), c["model"].lower()))


def watch_rows(db: DB, cfg: dict, records: list[PanelRecord], updated_keys: set[str]) -> list[dict]:
    """關注機型總覽：每個關注項目一列。"""
    from .enrich import watch_items, watch_matches
    rows = []
    for item in watch_items(cfg):
        recs = [r for r in records if watch_matches(item, r.spec) and r.confidence.rank >= Confidence.C.rank]
        if not recs:
            checked = db.watch_last_checked(item["name"])
            rows.append({"name": item["name"], "models": "", "screen": "", "score": "—", "score_pct": 0, "updated": False,
                         "status": "查無規格頁，下次回查再試" if checked else "排隊建檔中"})
            continue
        best = max(recs, key=lambda r: analytics.completeness(r)[0])
        k, t = analytics.completeness(best)
        models = list(dict.fromkeys(r.spec.model for r in recs if r.spec.model))
        card = panel_card(best)
        rows.append({"name": item["name"], "models": "、".join(models[:4]) + (f" 等 {len(models)} 款" if len(models) > 4 else ""),
                     "screen": card["screen"], "score": f"{k}/{t}", "score_pct": round(k / t * 100),
                     "updated": any(r.key in updated_keys for r in recs), "status": "",
                     "options": len(recs), "source_url": best.source_url})
    return rows


def backfill_status(db: DB, cfg: dict, added: int) -> dict | None:
    from . import backfill
    if not (cfg.get("backfill") or {}).get("enabled"):
        return None
    p = backfill.progress(db, cfg)
    return {"added": added, "quarters": p["quarters"], "done": p["done"], "total": p["total"],
            "start_on": cfg["backfill"].get("start_on")}


def ctx_label(weekly: bool) -> str:
    return "本週" if weekly else "今日"


def build_context(events: list[dict], db: DB, cfg: dict, today: datetime) -> dict:
    weekly = (cfg.get("schedule") or {}).get("cadence", "daily") == "weekly"
    period_start = (today - timedelta(days=7 if weekly else 1)).replace(hour=0, minute=0, second=0, microsecond=0)

    # 同一款機型在本期可能有多個事件（新增後又補規格），合併成一張卡片
    priority = {"rumor_confirmed": 3, "new": 2, "updated": 1}
    grouped: dict[str, dict] = {}
    for e in (e for e in events if e["category"] == "panel"):
        rec = _current_record(db, e)
        g = grouped.setdefault(rec.key, {"rec": rec, "kind": e["kind"], "changes": {}, "watch_seed": False, "backfill": False})
        if priority.get(e["kind"], 0) > priority.get(g["kind"], 0):
            g["kind"] = e["kind"]
        if e["payload"].get("tag") == "backfill":
            g["backfill"] = True  # 回溯的歷史資料：只進季度趨勢與 Excel，不出現在本期卡片
        if e["kind"] == "new" and e["payload"].get("tag") == "watch":
            g["watch_seed"] = True  # 關注清單主動建檔的既有機型，不算本期新機
        g["changes"].update(e["payload"].get("changes") or {})
    backfilled = sum(1 for g in grouped.values() if g["backfill"])
    grouped = {k: g for k, g in grouped.items() if not g["backfill"]}
    all_records = db.all_panels()
    period_new = [{"rec": g["rec"]} for g in grouped.values()
                  if g["kind"] == "new" and not g["watch_seed"] and g["rec"].confidence != Confidence.D]
    standout_lines, tags = analytics.standouts(period_new, all_records, period_start)

    def card(g):
        return panel_card(g["rec"], g["changes"] if g["kind"] != "new" else None, tags.get(g["rec"].key))

    by_name = lambda c: (c["brand"].lower(), c["model"].lower())
    is_rumor = lambda g: g["rec"].confidence == Confidence.D
    new_panels = sorted((card(g) for g in grouped.values() if g["kind"] == "new" and not g["watch_seed"] and not is_rumor(g)), key=by_name)
    updates = sorted((card(g) for g in grouped.values() if g["kind"] == "updated" and not is_rumor(g)), key=by_name)
    confirmed = sorted((card(g) for g in grouped.values() if g["kind"] == "rumor_confirmed"), key=by_name)
    rumor_panels = sorted((card(g) for g in grouped.values() if is_rumor(g)), key=by_name)
    ev_rumor = lambda e: e.get("confidence") == "D"

    platform = [e for e in events if e["category"] == "platform"]
    platform_rows = sorted(
        [{"vendor": e["payload"]["vendor"], "tech": e["payload"]["tech_name"],
          "category": PLATFORM_CATEGORY.get(e["payload"]["category"], ""),
          "summary": e["payload"]["summary"], "affected": "、".join(FIELD_LABELS.get(x, x) for x in e["payload"].get("affected_specs", [])),
          "confidence": Confidence(e["confidence"]).label if e.get("confidence") else "",
          "conf_code": e.get("confidence") or "C",
          "source_name": e["source_name"], "source_url": e["source_url"], "rumor": ev_rumor(e) or e["kind"] == "rumor"}
         for e in platform],
        key=lambda r: VENDOR_ORDER.index(r["vendor"]) if r["vendor"] in VENDOR_ORDER else 99)
    fab_rows = [{"maker": canonical_maker(e["payload"]["maker"]), "site": e["payload"].get("fab_site") or "",
                 "gen": e["payload"].get("generation") or "", "tech": e["payload"].get("technology") or "",
                 "status": FAB_STATUS.get(e["kind"], e["kind"]), "milestone": e["payload"].get("milestone") or "",
                 "capacity": e["payload"].get("capacity_sheets_per_month"),
                 "summary": e["payload"]["summary"], "confidence": Confidence(e["confidence"]).label if e.get("confidence") else "",
                 "source_name": e["source_name"], "source_url": e["source_url"], "rumor": ev_rumor(e)}
                for e in events if e["category"] == "fab"]
    industry = [{"note": e["payload"]["note"], "source_name": e["source_name"], "source_url": e["source_url"]}
                for e in events if e["category"] == "industry"]

    ctx = {
        "date": today.strftime("%Y-%m-%d"), "weekday": "一二三四五六日"[today.weekday()],
        "new_panels": new_panels, "updates": updates, "confirmed": confirmed,
        "rumor_panels": rumor_panels,
        "platform": [r for r in platform_rows if not r["rumor"]],
        "platform_rumors": [r for r in platform_rows if r["rumor"]],
        "fabs": [r for r in fab_rows if not r["rumor"]], "fab_rumors": [r for r in fab_rows if r["rumor"]],
        "industry": industry,
        "kpis": analytics.kpis(new_panels),
        "new_options": len(new_panels),
        "standouts": standout_lines,
        "wow": analytics.week_over_week(all_records, period_start, 7 if weekly else 1),
        "mentions": analytics.tech_mentions(all_records, period_start),
        "matrix": analytics.brand_maker_matrix(all_records),
        "rumor_tracker": analytics.rumor_tracker(all_records, today),
        "to_fill": analytics.to_fill(all_records, today),
        "platform_matrix": ihv.platform_matrix(all_records, today),
        "quarterly": analytics.quarterly_trends(all_records),
        "backfill": backfill_status(db, cfg, backfilled),
        "watch": watch_rows(db, cfg, all_records, set(grouped)),
        "watch_cmp": watch_comparison(db, cfg, all_records, set(grouped)) if cfg.get("watchlist") else None,
    }
    # 卡片依機型合併（螢幕選配列在同一張卡片內）
    for k in ("new_panels", "updates", "confirmed", "rumor_panels"):
        ctx[k + "_flat"] = ctx[k]
        ctx[k] = group_by_model(ctx[k])
    ctx["panel_known"] = sum(1 for c in ctx["new_panels"] if c["panel_known"])
    ctx["compare_new"] = comparison(ctx["new_panels"])
    if ctx["kpis"]:
        ctx["kpis"][0] = {"label": "新機型", "value": str(len(ctx["new_panels"])),
                          "note": f"共 {len(new_panels)} 種螢幕選配"}
        if ctx["new_panels"]:  # 完整度以每款機型資料最完整的選配計算，與卡片一致
            avg = sum(c["score_pct"] for c in ctx["new_panels"]) / len(ctx["new_panels"])
            for k in ctx["kpis"]:
                if k["label"] == "平均規格完整度":
                    k.update({"label": f"{ctx_label(weekly)}新機完整度", "value": f"{avg:.0f}%", "note": "剛發表，評測多半還沒出"})
        mature, n_mature = analytics.mature_completeness(all_records, today)
        ctx["kpis"].append({"label": "上市滿 30 天完整度", "value": f"{mature:.0%}" if mature is not None else "—",
                            "note": f"{n_mature} 款｜目標 75%" if mature is not None else "資料累積中｜目標 75%"})
    ctx["ongoing"] = ongoing(db, cfg, ctx)
    ctx["weekly"] = weekly
    ctx["title"] = "筆電顯示器情報週報" if weekly else "筆電顯示器情報日報"
    ctx["period_label"] = "本週" if weekly else "今日"
    ctx["prev_label"] = "上週" if weekly else "昨日"
    if weekly:
        ctx["period"] = f"資料期間：{(today - timedelta(days=6)).strftime('%m/%d')}–{today.strftime('%m/%d')}（過去 7 天）"
    else:
        ctx["period"] = "資料期間：前一日報告後至今"
    ctx["trend"] = weekly_trend(db) if (weekly or today.weekday() == 0) else None
    n_rumor = len(rumor_panels) + len(ctx["platform_rumors"]) + len(ctx["fab_rumors"])
    ctx["subject"] = (f"{cfg['report'].get('subject_prefix', '[顯示器情報]')} {ctx['date']}｜"
                      f"新機 {len(ctx['new_panels'])}・更新 {len(ctx['updates'])}・平台 {len(ctx['platform'])}・"
                      f"產線 {len(ctx['fabs'])}・傳聞 {n_rumor}")
    try:
        ctx["health"] = health.build(db, cfg, today, all_records, 7 if weekly else 1)
    except Exception as exc:  # 健康檢查失敗不影響寄報告
        log.warning("健康檢查失敗：%s", exc)
        ctx["health"] = None
    return ctx


def ongoing(db: DB, cfg: dict, ctx: dict) -> list[dict]:
    """今天沒有新消息的類別，列出近期仍在討論的主題。"""
    days = cfg["report"].get("ongoing_days", 30)
    recent = [e for e in db.recent_events(days) if e.get("reported_at")]
    out = []
    checks = [("新機型與面板", "panel", not ctx["new_panels"]),
              ("平台技術", "platform", not ctx["platform"]),
              ("面板廠產線", "fab", not ctx["fabs"])]
    for label, cat, empty in checks:
        if not empty:
            continue
        items = []
        for e in [e for e in recent if e["category"] == cat][:4]:
            p = e["payload"]
            if cat == "panel":
                s = p["record"]["spec"]
                title = f"{s.get('brand') or ''} {s.get('model') or ''} {s.get('panel_tech') or ''}".strip()
            elif cat == "platform":
                title = f"{p['vendor']} {p['tech_name']}：{p['summary']}"
            else:
                title = f"{p['maker']} {p.get('generation') or ''} {p.get('technology') or ''}：{p['summary']}"
            items.append({"title": title, "date": e["created_at"][:10], "source_url": e["source_url"]})
        out.append({"label": label, "entries": items})
    return out


def weekly_trend(db: DB) -> dict | None:
    """每週一：近 30 天新增紀錄的面板技術、刷新率分布。"""
    from datetime import timedelta, timezone
    since = datetime.now(timezone.utc) - timedelta(days=30)
    recs = [r for r in db.all_panels() if r.first_seen >= since and r.confidence.rank >= Confidence.C.rank]
    if len(recs) < 5:
        return None
    tech = Counter("OLED" if r.spec.panel_tech and "oled" in r.spec.panel_tech.lower() else
                   ("LCD" if r.spec.panel_tech else "未知") for r in recs)
    refresh = Counter(f"{r.spec.refresh_hz:g}Hz" for r in recs if r.spec.refresh_hz)
    makers = Counter(r.maker for r in recs if r.maker)
    return {"count": len(recs), "tech": tech.most_common(), "refresh": refresh.most_common(5),
            "makers": makers.most_common(5)}


def highlights(ctx: dict, llm=None) -> list[str]:
    """重點摘要：有 AI 就請 AI 寫，失敗就用規則產生。"""
    facts = {
        "新機": [f"{r['brand']} {r['model']} {r['size']} {r['tech']} {r['res']} {r['refresh']} {r['panel_src']}" for r in ctx["new_panels"][:15]],
        "規格更新": [f"{r['brand']} {r['model']}：{r['changes']}" for r in ctx["updates"][:10]],
        "傳聞證實": [f"{r['brand']} {r['model']}" for r in ctx["confirmed"]],
        "平台": [f"{r['vendor']} {r['tech']}：{r['summary']}" for r in ctx["platform"][:10]],
        "產線": [f"{r['maker']} {r['gen']} {r['status']}：{r['summary']}" for r in ctx["fabs"][:10]],
        "傳聞": [f"{r['brand']} {r['model']} {r['tech']} {r['refresh']}" for r in ctx["rumor_panels"][:10]],
        "關鍵數字": [f"{k['label']}：{k['value']}（{k['note']}）" for k in ctx.get("kpis", [])],
        "亮點": ctx.get("standouts", []),
        "與上期比較": [f"{w['metric']}：{w['prev']} → {w['now']}" for w in (ctx.get("wow") or [])],
    }
    if llm is not None and any(facts.values()):
        try:
            data = llm.complete_json(
                f"你是筆電顯示器市場分析師。根據提供的{ctx.get('period_label', '今日')}資料，用繁體中文寫 3 到 5 句重點，最重要的放第一句。"
                "只能使用提供的資料，不可加入資料以外的資訊。傳聞必須註明是傳聞。"
                "最後一句寫一個值得關注的市場或技術觀察（例如哪類規格變多、哪家面板廠出現較多），但仍只能根據資料。"
                '只輸出 JSON：{"highlights": ["...", "..."]}',
                json.dumps(facts, ensure_ascii=False))
            hl = [h for h in data.get("highlights", []) if isinstance(h, str) and h.strip()]
            if hl:
                return hl[:5]
        except Exception as exc:
            log.warning("AI 重點摘要失敗，改用規則摘要：%s", exc)
    out = []
    if ctx["new_panels"]:
        techs = Counter(r["tech"] or "未標示" for r in ctx["new_panels"])
        out.append(f"新增 {len(ctx['new_panels'])} 款機型面板紀錄，面板技術分布："
                   + "、".join(f"{k} {v} 款" for k, v in techs.most_common()) + "。")
    if ctx["updates"]:
        out.append(f"{len(ctx['updates'])} 款既有機型補上或修正規格。")
    if ctx["confirmed"]:
        out.append(f"{len(ctx['confirmed'])} 則先前傳聞獲得證實。")
    if ctx["platform"]:
        out.append(f"平台廠動態 {len(ctx['platform'])} 則：" + "、".join(sorted({r['vendor'] for r in ctx['platform']})) + "。")
    if ctx["fabs"]:
        out.append(f"面板廠產線動態 {len(ctx['fabs'])} 則。")
    return out or [f"{ctx.get('period_label', '今日')}沒有新的官方或媒體消息，下方列出近期仍在討論的主題。"]


def render_html(ctx: dict) -> str:
    env = Environment(loader=FileSystemLoader(TEMPLATES), autoescape=select_autoescape(["html", "j2"]))
    return env.get_template("report.html.j2").render(**ctx)


def write_excel(ctx: dict, db: DB, path: Path) -> Path:
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    fields = [f for f in FIELD_LABELS]

    def sheet(ws, headers, rows):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(10, max(len(str(c.value or "")) for c in col) + 2))

    write_compare_sheet(wb.active, ctx)
    ws = wb.create_sheet(f"{ctx.get('period_label', '今日')}新增與更新")
    rows = []
    for kind, items in [("新機", ctx.get("new_panels_flat", ctx["new_panels"])), ("更新", ctx.get("updates_flat", ctx["updates"])),
                        ("傳聞證實", ctx.get("confirmed_flat", ctx["confirmed"])), ("傳聞", ctx.get("rumor_panels_flat", ctx["rumor_panels"]))]:
        for r in items:
            s, rec = r["spec"], r["record"]
            rows.append([kind, r["confidence"]] + [_fmt(s.get(f)) for f in fields]
                        + [rec.get("ppi"), rec.get("aspect_ratio"), rec.get("density_class"), r["changes"], r["source_name"], r["source_url"]])
    sheet(ws, ["類型", "可信度"] + [FIELD_LABELS[f] for f in fields] + ["PPI", "長寬比", "像素密度", "變更", "來源", "網址"], rows)
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("全部面板紀錄")
    sheet(ws2, ["可信度"] + [FIELD_LABELS[f] for f in fields]
          + ["表面代碼", "面板廠（含推算）", "PPI", "長寬比", "像素密度", "首次出現", "最後更新", "來源", "網址", "查核標記"],
          [[rec.confidence.label] + [_fmt(getattr(rec.spec, f)) for f in fields]
           + [derive.surface_code(rec.spec.surface) or "", (f"{rec.maker}（推算）" if rec.panel_maker_inferred else rec.maker) or "",
              rec.ppi, rec.aspect_ratio, rec.density_class, rec.first_seen.strftime("%Y-%m-%d"),
              rec.last_updated.strftime("%Y-%m-%d"), rec.source_name, rec.source_url, "；".join(rec.flags)]
           for rec in db.all_panels()])
    ws2.auto_filter.ref = ws2.dimensions

    m = ctx.get("matrix")
    if m:
        ws4 = wb.create_sheet("品牌×面板廠")
        sheet(ws4, ["品牌"] + m["makers"], [[r["brand"]] + r["cells"] for r in m["rows"]])
    if ctx.get("watch"):
        from .enrich import watch_items, watch_matches
        ws7 = wb.create_sheet("關注機型")
        items = watch_items(ctx.get("_cfg") or {})
        rows7 = []
        for rec in db.all_panels():
            item = next((i for i in items if watch_matches(i, rec.spec)), None)
            if item and rec.confidence.rank >= Confidence.C.rank:
                rows7.append([item["name"], rec.confidence.label, "%d/%d" % analytics.completeness(rec)]
                             + [_fmt(getattr(rec.spec, f)) for f in fields]
                             + [rec.ppi, rec.last_updated.strftime("%Y-%m-%d"), rec.source_name, rec.source_url])
        sheet(ws7, ["關注項目", "可信度", "完整度"] + [FIELD_LABELS[f] for f in fields] + ["PPI", "最後更新", "來源", "網址"],
              sorted(rows7, key=lambda r: (r[0], r[4] or "")))
        ws7.auto_filter.ref = ws7.dimensions
    qt = ctx.get("quarterly")
    if qt and qt["dims"]:
        ws8 = wb.create_sheet("季度趨勢")
        rows8 = []
        for d in qt["dims"]:
            for row in d["rows"]:
                for b in d["buckets"]:
                    share = row["share"].get(b)
                    rows8.append([d["name"], row["quarter"], row["n"], b, row["counts"].get(b, 0),
                                  round(share * 100, 1) if share is not None else None,
                                  d["shift"].get(b) if row is d["rows"][-1] else None])
        sheet(ws8, ["項目", "季度", "有標示的機型數", "分類", "款數", "占比 %", "季變化（百分點）"], rows8)
        ws8.auto_filter.ref = ws8.dimensions
    ws6 = wb.create_sheet("平台適配")
    names = [i["name"] for i in ihv.evaluate(PanelSpec())]
    sheet(ws6, ["品牌", "機型", "螢幕", "處理器", "處理器平台", "顯示卡", "獨顯廠商", "顯示切換"] + names + ["說明"],
          [[rec.spec.brand, rec.spec.model, f"{_fmt(rec.spec.size_inch)} {rec.spec.panel_tech or ''} {_fmt(rec.spec.refresh_hz)}Hz".strip(),
            rec.spec.cpu_platform or "", ihv.cpu_vendor(rec.spec) or "", rec.spec.gpu or "", ihv.gpu_vendor(rec.spec) or "",
            rec.spec.gpu_switching or ""]
           + [i["status"] for i in ihv.evaluate(rec.spec)]
           + ["；".join(f"{i['name']}：{i['note']}" for i in ihv.evaluate(rec.spec) if i["note"])]
           for rec in db.all_panels() if rec.confidence.rank >= Confidence.C.rank])
    ws6.auto_filter.ref = ws6.dimensions
    if ctx.get("mentions"):
        ws5 = wb.create_sheet("技術提及率")
        sheet(ws5, ["技術", "說明", f"{ctx.get('period_label', '本期')}款數", f"{ctx.get('period_label', '本期')}提及率", "前 30 天提及率"],
              [[r["feature"], r["desc"], r["count"], r["share"], r["base_share"]] for r in ctx["mentions"]])

    hc = ctx.get("health")
    if hc:
        wsh = wb.create_sheet("資料查核")
        rows_h = [[sec, c["icon"] + " " + {"ok": "正常", "warn": "注意", "bad": "需處理"}[c["level"]], c["name"], c["detail"], it]
                  for sec, key in (("運作", "ops"), ("資料", "quality")) for c in hc[key] for it in (c["notes"] or [""])]
        rows_h += [["抽樣稽核", a["confidence"], a["model"], f"{ln['field']}：{_fmt(ln['value'])}", f"原文：「{ln['quote']}」　{a['source_url']}"]
                   for a in hc["audit"] for ln in a["lines"]]
        sheet(wsh, ["類別", "狀態", "檢查項目", "結果", "明細"], rows_h)
        wsh.column_dimensions["D"].width = 50
        wsh.column_dimensions["E"].width = 80
    ws3 = wb.create_sheet("平台與產線")
    sheet(ws3, ["類別", "廠商", "項目", "分類／狀態", "摘要", "可信度", "來源", "網址"],
          [["平台", r["vendor"], r["tech"], r["category"], r["summary"], r["confidence"], r["source_name"], r["source_url"]]
           for r in ctx["platform"] + ctx["platform_rumors"]]
          + [["產線", r["maker"], f"{r['site']} {r['gen']} {r['tech']}".strip(), r["status"], r["summary"], r["confidence"], r["source_name"], r["source_url"]]
             for r in ctx["fabs"] + ctx["fab_rumors"]])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def build_report(db: DB, cfg: dict, llm=None) -> tuple[dict, str, Path, list[int]]:
    """產生報告。回傳（內容, HTML, Excel 路徑, 已報告的事件 id）。"""
    tz = ZoneInfo(cfg["schedule"].get("timezone", "Asia/Taipei"))
    today = datetime.now(tz)
    events = db.unreported_events()
    ctx = build_context(events, db, cfg, today)
    ctx["_cfg"] = cfg
    ctx["highlights"] = highlights(ctx, llm)
    ctx["glossary"] = []
    html = render_html(ctx)
    ctx["glossary"] = analytics.glossary_for(html)
    html = render_html(ctx)
    out = Path(cfg["report"]["output_dir"])
    out.mkdir(parents=True, exist_ok=True)
    (out / f"report-{ctx['date']}.html").write_text(html, encoding="utf-8")
    xlsx = write_excel(ctx, db, out / f"display-intel-{ctx['date']}.xlsx")
    return ctx, html, xlsx, [e["id"] for e in events]


# ---------- 比較表 ----------

SEGMENT_ORDER = ["13 吋以下", "14 吋級", "15.6 吋級", "16 吋級", "17–18 吋", "AIO", "尺寸未標示"]
# 欄位 → 數值取法、最佳方向（max / min）
BEST_RULES = {
    "refresh": (lambda s, r: s.refresh_hz, max), "response": (lambda s, r: s.response_ms, min),
    "sdr": (lambda s, r: s.sdr_nits, max), "hdr": (lambda s, r: s.hdr_peak_nits, max),
    "p3": (lambda s, r: s.dci_p3_pct, max), "ppi": (lambda s, r: r.ppi, max),
}


def _segment(rec: PanelRecord) -> str:
    return analytics._size_bucket(rec) or "尺寸未標示"


def compare_row(rec: PanelRecord, label: str = "", option: str = "", extra: dict | None = None) -> dict:
    s = rec.spec
    code = derive.surface_code(s.surface)
    gamut = [f"P3 {s.dci_p3_pct:g}%" if s.dci_p3_pct else "", f"sRGB {s.srgb_pct:g}%" if s.srgb_pct else "",
             f"Adobe {s.adobe_rgb_pct:g}%" if s.adobe_rgb_pct else "", f"NTSC {s.ntsc_pct:g}%" if s.ntsc_pct else ""]
    maker = rec.maker
    k, t = analytics.completeness(rec)
    return {
        "rec": rec, "label": label, "option": option, "segment": _segment(rec),
        "size": f'{s.size_inch:g}"' if s.size_inch else (f'{rec.size_inferred:g}"*' if rec.size_inferred else "—"),
        "tech": s.panel_tech or "—",
        "res": f"{s.res_w}×{s.res_h}" if s.res_w and s.res_h else "—", "aspect": rec.aspect_ratio if rec.aspect_ratio and len(rec.aspect_ratio) <= 5 else "",
        "refresh": f"{s.refresh_hz:g}Hz" if s.refresh_hz else "—",
        "response": f"{s.response_ms:g}ms" if s.response_ms else "",
        "sdr": f"{s.sdr_nits:g}" if s.sdr_nits else "—", "hdr": f"{s.hdr_peak_nits:g}" if s.hdr_peak_nits else "",
        "gamut": [g for g in gamut if g] or ["—"],
        "surface": code or ((s.surface[:14] + ("…" if len(s.surface) > 14 else "")) if s.surface else "—"),
        "maker": (f"{maker}*" if rec.panel_maker_inferred else maker) if maker else "—",
        "panel_model": s.panel_model or "",
        "touch": "觸控" if s.touch else "",
        "score": f"{k}/{t}", "score_pct": round(k / t * 100),
        "best": set(), "source_url": rec.source_url, **(extra or {}),
    }


def mark_best(rows: list[dict]):
    """在同一組裡標出各欄最佳值（同分都標；少於 2 個有效值不標）。"""
    for key, (get, pick) in BEST_RULES.items():
        vals = [(get(r["rec"].spec, r["rec"]), r) for r in rows]
        vals = [(v, r) for v, r in vals if v]
        if len(vals) < 2:
            continue
        best = pick(v for v, _ in vals)
        for v, r in vals:
            if v == best:
                r["best"].add(key)


def comparison(groups: list[dict]) -> list[dict]:
    """把（已依機型合併的）卡片轉成比較表，依尺寸級距分組。"""
    segs: dict[str, list] = {}
    for c in groups:
        opts = c.get("options") or [c]
        for i, o in enumerate(opts):
            row = compare_row(o["rec"], f"{c['brand']} {c['model']}" if i == 0 else "",
                              f"選配 {i + 1}" if len(opts) > 1 else "", {"tags": c.get("tags", []) if i == 0 else []})
            segs.setdefault(row["segment"], []).append(row)
    out = []
    for name in sorted(segs, key=lambda x: SEGMENT_ORDER.index(x) if x in SEGMENT_ORDER else 99):
        mark_best(segs[name])
        _stripe(segs[name])
        out.append({"name": name, "rows": segs[name]})
    return out


def _stripe(rows: list[dict]):
    """同一款機型的選配同底色，機型之間交替，方便閱讀。"""
    on = True
    for r in rows:
        if r["label"]:
            on = not on
        r["stripe"] = on


def watch_comparison(db: DB, cfg: dict, records: list[PanelRecord], updated_keys: set[str]) -> dict:
    """關注機型比較表：每個關注項目取資料最完整的選配；尚未建檔的列在 pending。"""
    from .enrich import watch_items, watch_matches
    rows, pending = [], []
    for item in watch_items(cfg):
        recs = [r for r in records if watch_matches(item, r.spec) and r.confidence.rank >= Confidence.C.rank]
        if not recs:
            pending.append({"name": item["name"], "status": "查無規格頁，下次回查" if db.watch_last_checked(item["name"]) else "排隊建檔中"})
            continue
        best = max(recs, key=lambda r: analytics.completeness(r)[0])
        models = list(dict.fromkeys(r.spec.model for r in recs if r.spec.model))
        rows.append(compare_row(best, item["name"], f"{len(recs)} 種選配" if len(recs) > 1 else "",
                                {"updated": any(r.key in updated_keys for r in recs), "models": "、".join(models[:3])}))
    mark_best(rows)
    _stripe(rows)
    return {"rows": rows, "pending": pending}


# ---------- Excel 比較表（第一個工作表） ----------

def _j(v) -> str:
    if v is None or v == "" or v == []:
        return ""
    if isinstance(v, list):
        return "、".join(map(str, v))
    if isinstance(v, bool):
        return "有" if v else "無"
    return _fmt(v)


# (群組, 欄名, 取值, 最佳方向 max/min/None, 欄寬)
COMPARE_COLUMNS = [
    ("基本", "機型", lambda r: r["label"], None, 26),
    ("基本", "選配", lambda r: r["option"], None, 9),
    ("基本", "尺寸", lambda r: r["rec"].spec.size_inch or r["rec"].size_inferred, None, 7),
    ("基本", "面板", lambda r: r["rec"].spec.panel_tech, None, 9),
    ("基本", "解析度", lambda r: r["res"] if r["res"] != "—" else "", None, 12),
    ("基本", "長寬比", lambda r: r["aspect"], None, 8),
    ("基本", "PPI", lambda r: r["rec"].ppi, max, 7),
    ("動態", "刷新率 Hz", lambda r: r["rec"].spec.refresh_hz, max, 9),
    ("動態", "反應 ms", lambda r: r["rec"].spec.response_ms, min, 8),
    ("動態", "VRR", lambda r: _j(r["rec"].spec.vrr_range) or _j(r["rec"].spec.vrr_tech), None, 12),
    ("亮度與色彩", "SDR nits", lambda r: r["rec"].spec.sdr_nits, max, 9),
    ("亮度與色彩", "HDR nits", lambda r: r["rec"].spec.hdr_peak_nits, max, 9),
    ("亮度與色彩", "DCI-P3 %", lambda r: r["rec"].spec.dci_p3_pct, max, 8),
    ("亮度與色彩", "sRGB %", lambda r: r["rec"].spec.srgb_pct, max, 8),
    ("亮度與色彩", "Adobe %", lambda r: r["rec"].spec.adobe_rgb_pct, max, 8),
    ("亮度與色彩", "NTSC %", lambda r: r["rec"].spec.ntsc_pct, max, 8),
    ("亮度與色彩", "ΔE", lambda r: r["rec"].spec.delta_e, min, 6),
    ("亮度與色彩", "色深", lambda r: r["rec"].spec.color_depth_desc or (f"{r['rec'].spec.color_depth_bit}-bit" if r["rec"].spec.color_depth_bit else ""), None, 10),
    ("表面與光學", "表面", lambda r: derive.surface_code(r["rec"].spec.surface) or "", None, 8),
    ("表面與光學", "表面原文", lambda r: r["rec"].spec.surface, None, 16),
    ("表面與光學", "反射率 %", lambda r: r["rec"].spec.reflectance_pct, min, 8),
    ("表面與光學", "均勻度 %", lambda r: r["rec"].spec.uniformity_pct, max, 8),
    ("表面與光學", "可視角", lambda r: r["rec"].spec.viewing_angle, None, 10),
    ("結構", "背板", lambda r: r["rec"].spec.backplane, None, 8),
    ("結構", "背光／調光", lambda r: " / ".join(x for x in [r["rec"].spec.backlight, r["rec"].spec.dimming] if x), None, 14),
    ("結構", "觸控", lambda r: (("觸控 " + (r["rec"].spec.touch_structure or "")).strip() if r["rec"].spec.touch
                              else ("非觸控" if r["rec"].spec.touch is False else "")), None, 12),
    ("機構", "屏佔比 %", lambda r: r["rec"].spec.screen_to_body_pct, max, 8),
    ("機構", "面板厚度 mm", lambda r: r["rec"].spec.panel_thickness_mm, min, 9),
    ("機構", "邊框", lambda r: r["rec"].spec.bezel, None, 12),
    ("機構", "鏡頭", lambda r: r["rec"].spec.camera_cutout, None, 10),
    ("介面", "eDP", lambda r: r["rec"].spec.edp_version, None, 9),
    ("介面", "VRR 最低 Hz", lambda r: r["rec"].spec.vrr_min_hz, min, 9),
    ("介面", "面板功耗 W", lambda r: r["rec"].spec.panel_power_w, min, 9),
    ("介面", "省電技術", lambda r: _j(r["rec"].spec.low_power_tech) or _j(r["rec"].spec.display_power_tech), None, 14),
    ("面板來源", "面板廠", lambda r: r["maker"] if r["maker"] != "—" else "", None, 14),
    ("面板來源", "面板型號", lambda r: r["panel_model"], None, 14),
    ("面板來源", "完整度", lambda r: r["score"], None, 7),
    ("面板來源", "可信度", lambda r: r["rec"].confidence.label, None, 9),
    ("面板來源", "網址", lambda r: r["source_url"], None, 30),
]
GROUP_COLORS = {"基本": "1F4E79", "動態": "2E75B6", "亮度與色彩": "C55A11", "表面與光學": "548235",
                "結構": "7030A0", "機構": "806000", "介面": "2F5597", "面板來源": "404040"}


def write_compare_sheet(ws, ctx: dict):
    """比較表：兩層表頭（欄位群組）、依尺寸分段、同段最佳值標綠、機型交替底色。"""
    ws.title = "比較表"
    white = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    best_fill, best_font = PatternFill("solid", fgColor="C6EFCE"), Font(bold=True, color="006100")
    stripe_fill = PatternFill("solid", fgColor="F2F2F2")
    sec_fill, seg_fill = PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="BFBFBF")
    n = len(COMPARE_COLUMNS)

    # 第 1–2 列：群組表頭＋欄名
    col = 1
    for i, (grp, name, _, _, width) in enumerate(COMPARE_COLUMNS, 1):
        fill = PatternFill("solid", fgColor=GROUP_COLORS.get(grp, "404040"))
        for row, val in ((1, grp), (2, name)):
            c = ws.cell(row=row, column=i, value=val)
            c.font, c.fill, c.alignment = white, fill, center
        ws.column_dimensions[get_column_letter(i)].width = width
    start = 1
    for i in range(1, n + 2):
        if i > n or COMPARE_COLUMNS[i - 1][0] != COMPARE_COLUMNS[start - 1][0]:
            if i - 1 > start:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i - 1)
            start = i
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "C3"

    def section(title, fill):
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12 if fill is sec_fill else 11)
        for i in range(1, n + 1):
            ws.cell(row=r, column=i).fill = fill
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)

    def best_cols(rows):
        out = {}
        for idx, (_, _, get, pick, _) in enumerate(COMPARE_COLUMNS):
            if pick is None:
                continue
            vals = [(get(r), id(r)) for r in rows]
            vals = [(v, k) for v, k in vals if isinstance(v, (int, float)) and v]
            if len(vals) < 2:
                continue
            b = pick(v for v, _ in vals)
            out[idx] = {k for v, k in vals if v == b}
        return out

    def table(rows):
        best = best_cols(rows)
        for r in rows:
            ws.append([("" if (v := get(r)) is None else v) for _, _, get, _, _ in COMPARE_COLUMNS])
            rr = ws.max_row
            for idx in range(n):
                c = ws.cell(row=rr, column=idx + 1)
                c.border = Border(bottom=thin)
                c.alignment = Alignment(vertical="center", wrap_text=COMPARE_COLUMNS[idx][4] >= 12)
                if r.get("stripe"):
                    c.fill = stripe_fill
                if id(r) in best.get(idx, ()):
                    c.fill, c.font = best_fill, best_font
            if r["label"]:
                ws.cell(row=rr, column=1).font = Font(bold=True)
            url = r.get("source_url")
            if url:
                c = ws.cell(row=rr, column=n)
                c.hyperlink, c.font = url, Font(color="0563C1", underline="single")

    period = ctx.get("period_label", "本期")
    segs = ctx.get("compare_new") or []
    section(f"{period}新機（依尺寸分段；綠底＝同段最佳；* 為由面板型號推算）", sec_fill)
    if not segs:
        section("（本期沒有新機）", PatternFill(fill_type=None))
    for seg in segs:
        section(f"▸ {seg['name']}（{len({r['rec'].spec.model for r in seg['rows']})} 款）", seg_fill)
        table(seg["rows"])
    w = ctx.get("watch_cmp") or {}
    if w.get("rows") or w.get("pending"):
        ws.append([])
        section("關注機型（每個項目取資料最完整的選配；綠底＝關注清單中最佳）", sec_fill)
        table(w.get("rows", []))
        if w.get("pending"):
            ws.append(["尚未建檔：" + "、".join(f"{p['name']}（{p['status']}）" for p in w["pending"])])
            ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="7F7F7F")
    ws.append([])
    ws.append(["說明：空白＝文章沒寫（不是不支援）；表面代碼 AG 霧面、AGLR 抗眩光低反射、AR 抗反射、BV 亮面。"])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="7F7F7F")
    return ws
