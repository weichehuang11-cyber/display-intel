"""人工更正：config.yaml 的 corrections 清單，最高優先，網路資料不能覆蓋。

每次蒐集與每次產生週報前都會套用（重複套用不會重複產生更新事件）。
- 數值有實際改變時產生 updated 事件（tag="manual"），週報「規格更新」會列出來確認已生效
- 設成 null（空白）＝確認這個數值是錯的，保持空白，之後的新聞也不能再填
- 找不到對應機型的更正，健康檢查會提醒
"""

from __future__ import annotations

import logging

from .db import DB, utcnow
from .panel_decode import canonical_maker
from .schema import Confidence, FieldSource, PanelEvent, PanelRecord, PanelSpec

log = logging.getLogger(__name__)

IDENTITY = {"brand", "model"}


def items(cfg: dict) -> list[dict]:
    """全部人工更正：config.yaml 的 corrections ＋ corrections 資料夾裡的 Excel／CSV。"""
    out = [c for c in (cfg.get("corrections") or []) if isinstance(c, dict) and c.get("brand") and c.get("model")]
    return out + load_files(cfg)[0]


def _label(c: dict) -> str:
    extra = "、".join(f"{k}={v}" for k, v in (c.get("match") or {}).items())
    where = f"［{c['_origin']}］" if c.get("_origin") else ""
    return f"{where}{c['brand']} {c['model']}" + (f"（{extra}）" if extra else "")


# ---------- Excel／CSV 更正檔（corrections 資料夾）----------
# 表頭（括號內是說明，比對時忽略）→ 欄位。品牌、機型必填；空白＝不更動；填「清空」＝確認錯誤、保持空白
SHEET_COLUMNS = [
    ("品牌", "brand"), ("機型", "model"), ("比對尺寸（吋，選填）", "match_size"), ("比對解析度（選填，例 2560x1600）", "match_res"),
    ("面板廠", "panel_maker"), ("面板型號", "panel_model"), ("面板技術", "panel_tech"), ("尺寸（吋）", "size_inch"),
    ("解析度（例 2560x1600）", "res"), ("刷新率（Hz）", "refresh_hz"), ("反應時間（ms）", "response_ms"),
    ("SDR 亮度（nits）", "sdr_nits"), ("HDR 峰值（nits）", "hdr_peak_nits"), ("DCI-P3（%）", "dci_p3_pct"),
    ("sRGB（%）", "srgb_pct"), ("Adobe RGB（%）", "adobe_rgb_pct"), ("NTSC（%）", "ntsc_pct"), ("表面處理", "surface"),
    ("觸控（是／否）", "touch"), ("背板", "backplane"), ("色深", "color_depth_desc"), ("調光", "dimming"),
    ("PWM（Hz）", "pwm_hz"), ("起價", "price"), ("幣別（USD、TWD…）", "price_currency"),
    ("排除（是＝不列入報告）", "exclude"), ("來源說明", "note"),
]
_TRUE = {"是", "有", "支援", "y", "yes", "true", "✓", "v", "1"}
_FALSE = {"否", "無", "不支援", "n", "no", "false", "✗", "x", "0"}


def _head_key(h) -> str:
    import re
    return re.sub(r"[（(].*?[）)]|\s", "", str(h or "")).lower()


_HEAD_MAP = {_head_key(h): f for h, f in SHEET_COLUMNS}


def _res(v) -> tuple[int, int] | None:
    import re
    m = re.fullmatch(r"\s*(\d{3,4})\s*[x×*X]\s*(\d{3,4})\s*", str(v))
    return (int(m.group(1)), int(m.group(2))) if m else None


def _rows_from_file(path) -> list[list]:
    if path.suffix.lower() == ".csv":
        import csv
        for enc in ("utf-8-sig", "big5", "cp950"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    return [r for r in csv.reader(f)]
            except UnicodeDecodeError:
                continue
        return []
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["人工更正"] if "人工更正" in wb.sheetnames else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def parse_rows(rows: list[list], origin: str) -> tuple[list[dict], list[str]]:
    """把表格轉成更正項目。回傳（項目, 錯誤說明）。"""
    out, errors = [], []
    head_i = next((i for i, r in enumerate(rows[:20]) if {"品牌", "機型"} <= {_head_key(c) for c in r}), None)
    if head_i is None:
        return [], [f"［{origin}］找不到「品牌」「機型」表頭"]
    fields = [_HEAD_MAP.get(_head_key(c)) for c in rows[head_i]]
    unknown = [str(c) for c, f in zip(rows[head_i], fields) if c not in (None, "") and not f]
    if unknown:
        errors.append(f"［{origin}］看不懂的欄位（已略過）：{'、'.join(unknown)}")
    for n, row in enumerate(rows[head_i + 1:], start=head_i + 2):
        vals = {f: (str(v).strip() if v is not None else "") for f, v in zip(fields, row) if f}
        if not vals.get("brand") or not vals.get("model"):
            continue
        item = {"brand": vals["brand"], "model": vals["model"], "set": {}, "match": {}, "_origin": f"{origin} 第 {n} 列",
                "note": vals.get("note") or f"人工更正檔 {origin}"}
        if vals.get("match_size"):
            item["match"]["size_inch"] = vals["match_size"]
        if vals.get("match_res"):
            r = _res(vals["match_res"])
            if r:
                item["match"].update({"res_w": r[0], "res_h": r[1]})
            else:
                errors.append(f"［{item['_origin']}］比對解析度格式不對：{vals['match_res']}（例 2560x1600）")
        for f, v in vals.items():
            if f in ("brand", "model", "note", "match_size", "match_res") or v == "":
                continue
            if f == "exclude":
                if v.lower() in _TRUE:
                    item["exclude"] = True
                continue
            if v in ("清空", "刪除", "null"):
                if f == "res":
                    item["set"].update({"res_w": None, "res_h": None})
                else:
                    item["set"][f] = None
            elif f == "res":
                r = _res(v)
                if r:
                    item["set"].update({"res_w": r[0], "res_h": r[1]})
                else:
                    errors.append(f"［{item['_origin']}］解析度格式不對：{v}（例 2560x1600）")
            elif f == "touch":
                low = v.lower()
                if low in _TRUE or low in _FALSE:
                    item["set"]["touch"] = low in _TRUE
                else:
                    errors.append(f"［{item['_origin']}］觸控請填「是」或「否」：{v}")
            else:
                item["set"][f] = v
        if item["set"] or item.get("exclude"):
            out.append(item)
    return out, errors


def corrections_dir(cfg: dict):
    from .config import ROOT
    from pathlib import Path
    d = Path(cfg.get("corrections_dir") or "corrections")
    return d if d.is_absolute() else ROOT / d


def load_files(cfg: dict) -> tuple[list[dict], list[str]]:
    """讀取 corrections 資料夾裡所有 .xlsx／.csv（檔名開頭是 ~$ 的 Excel 暫存檔略過）。"""
    d = corrections_dir(cfg)
    if not d.is_dir():
        return [], []
    out, errors = [], []
    for path in sorted(d.iterdir()):
        if path.suffix.lower() not in (".xlsx", ".csv") or path.name.startswith(("~$", ".")):
            continue
        try:
            items_, errs = parse_rows(_rows_from_file(path), path.name)
        except Exception as exc:
            items_, errs = [], [f"［{path.name}］無法讀取：{type(exc).__name__}"]
        out += items_
        errors += errs
    return out, errors


def write_template(ws, rows: list[dict] | None = None):
    """在工作表寫入人工更正範本：說明列＋表頭＋（可選）預先填好品牌、機型、比對尺寸與解析度的列。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    ws.title = "人工更正"
    n = len(SHEET_COLUMNS)
    ws.cell(row=1, column=1, value="人工更正：你確認過的規格填在這裡，存檔後上傳到 GitHub 的 corrections 資料夾，下次蒐集或週報就會生效。").font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value="規則：品牌、機型必填；空白＝不更動；填「清空」＝確認原本的數值是錯的、保持空白；"
                                   "比對尺寸／解析度可指定只改某個螢幕選配（已預先填好）；來源說明會顯示在報告的出處。").font = Font(color="666666")
    for c in (1, 2):
        ws.merge_cells(start_row=c, start_column=1, end_row=c, end_column=n)
    head_fill, key_fill = PatternFill("solid", fgColor="1F3A5F"), PatternFill("solid", fgColor="EEF2F7")
    for j, (h, f) in enumerate(SHEET_COLUMNS, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = 18 if f in ("model", "note", "panel_model") else 12
    ws.row_dimensions[3].height = 34
    ws.freeze_panes = "C4"
    for i, r in enumerate(rows or [], start=4):
        for j, (_, f) in enumerate(SHEET_COLUMNS, 1):
            if f in r and r[f] not in (None, ""):
                ws.cell(row=i, column=j, value=r[f])
            if f in ("brand", "model", "match_size", "match_res"):
                ws.cell(row=i, column=j).fill = key_fill
    last = max(4 + len(rows or []) + 50, 60)
    col = {f: get_column_letter(j) for j, (_, f) in enumerate(SHEET_COLUMNS, 1)}
    for f, options in (("touch", "是,否,清空"), ("exclude", "是"), ("panel_tech", "IPS LCD,VA LCD,Mini LED LCD,OLED,Tandem OLED,清空"),
                       ("surface", "AG,AGLR,AR,BV,BV+AR,清空")):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{col[f]}4:{col[f]}{last}")
    return ws


def matches(c: dict, rec: PanelRecord) -> bool:
    from .verify import model_norm
    s = rec.spec
    if (s.brand or "").strip().lower() != str(c["brand"]).strip().lower():
        return False
    if model_norm(s) != model_norm(PanelSpec(brand=str(c["brand"]), model=str(c["model"]))):
        return False
    for k, v in (c.get("match") or {}).items():
        cur = getattr(s, k, None)
        if cur is None or (str(cur).lower() != str(v).lower() and not (isinstance(cur, (int, float)) and _num(v) == cur)):
            return False
    return True


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def validate(c: dict) -> tuple[dict, list[str]]:
    """檢查欄位名稱與型別；回傳（可套用的值, 錯誤說明）。"""
    errors, values = [], {}
    if c.get("exclude"):
        return {}, []
    for f, v in (c.get("set") or {}).items():
        if f not in PanelSpec.model_fields or f in IDENTITY or f == "evidence":
            errors.append(f"{_label(c)}：沒有「{f}」這個欄位")
            continue
        if v is None:
            values[f] = None
            continue
        try:
            values[f] = getattr(PanelSpec(**{f: v}), f)
        except Exception:
            errors.append(f"{_label(c)}：{f} 的值「{v}」格式不對")
            continue
        if f == "panel_maker":
            values[f] = canonical_maker(values[f]) or values[f]
    return values, errors


def apply(db: DB, cfg: dict) -> dict:
    """套用全部人工更正。回傳統計（matched / changed / unmatched / errors）。"""
    from .verify import _derived
    stats = {"items": 0, "matched": 0, "changed": 0, "excluded": 0, "unmatched": [], "errors": []}
    corr = items(cfg)
    records = db.all_panels()
    # 排除來源：這些網頁（例如多年前的舊文、舊特價文）帶進來的機型不列入報告
    for pat in cfg.get("exclude_sources") or []:
        for rec in records:
            doms = {fs.domain for fs in rec.field_sources.values() if fs.kind == "article"}
            if str(pat) in rec.source_url and len(doms) <= 1 and rec.confidence != Confidence.E:
                _exclude(rec, f"來源已排除：{pat}")
                db.put_panel(rec)
                stats["excluded"] += 1
    if not corr:
        return stats
    now = utcnow()
    for c in corr:
        stats["items"] += 1
        values, errors = validate(c)
        stats["errors"] += errors
        targets = [r for r in records if matches(c, r)]
        if not targets:
            stats["unmatched"].append(_label(c))
            continue
        note = str(c.get("note") or "").strip()
        if c.get("exclude"):
            for rec in targets:
                stats["matched"] += 1
                if rec.confidence != Confidence.E:
                    _exclude(rec, "人工排除" + (f"：{note}" if note else ""))
                    db.put_panel(rec)
                    stats["excluded"] += 1
            continue
        for rec in targets:
            stats["matched"] += 1
            data = rec.spec.model_dump()
            changes, touched = {}, False
            for f, v in values.items():
                if data.get(f) != v:
                    changes[f] = [data.get(f), v]
                    data[f] = v
                fs = rec.field_sources.get(f)
                if fs is None or fs.kind != "manual" or fs.source != note or f in changes:
                    touched = True
                    rec.field_sources[f] = FieldSource(source=note, conf=Confidence.A, at=now, kind="manual")
                    rec.corroborated.pop(f, None)
                    if v is None:
                        data["evidence"].pop(f, None)
                    else:
                        data["evidence"][f] = "人工更正" + (f"：{note}" if note else "")
            if not touched:
                continue  # 已套用過，什麼都不用做
            rec.spec = PanelSpec(**data)
            if changes:
                rec.last_updated = now
                rec.flags = [fl for fl in rec.flags if not any(f"：{f}（" in fl for f in changes)]  # 已人工裁決的不一致不再列
            rec = _derived(rec)
            db.put_panel(rec)
            if changes:
                stats["changed"] += 1
                ev = PanelEvent(kind="updated", record=rec, changes=changes)
                payload = ev.model_dump(mode="json") | {"tag": "manual"}
                db.add_event("panel", "updated", payload, Confidence.A.value, "", "人工更正" + (f"：{note}" if note else ""))
    if stats["changed"] or stats["unmatched"] or stats["errors"]:
        log.info("人工更正：%d 項，套用到 %d 筆，改變 %d 筆；找不到 %s；錯誤 %s", stats["items"], stats["matched"],
                 stats["changed"], stats["unmatched"] or "無", stats["errors"] or "無")
    return stats


def _exclude(rec: PanelRecord, why: str):
    """改成待確認（E）：不列入報告，資料保留；原可信度記在 flags，方便恢復。"""
    rec.flags.append(f"{why}（原可信度 {rec.confidence.value}，改為待確認、不列入報告）")
    rec.confidence = Confidence.E


def unmatched(records: list[PanelRecord], cfg: dict) -> tuple[list[str], list[str]]:
    """健康檢查用：找不到機型的更正、格式錯誤。"""
    miss, errs = [], list(load_files(cfg)[1])
    for c in items(cfg):
        errs += validate(c)[1]
        if not any(matches(c, r) for r in records):
            miss.append(_label(c))
    return miss, errs
